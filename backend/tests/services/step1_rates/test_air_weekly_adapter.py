import uuid

from openpyxl import Workbook

from app.services.step1_rates.sheet_builder.template_filler import fill_template
from app.services.step1_rates.adapters.air_weekly import AirWeeklyAdapter
from app.services.step1_rates.adapters.air import AirAdapter
from app.services.step1_rates.activator_mappers import to_air_freight_rate
from app.services.step1_rates.entities import Step1FileType


def _make(tmp_path):
    """做表生成的空运周表（严格按客户原件模板：无起运港列、无币种列）。"""
    rows = [{
        "origin": "PVG", "destination": "NRT", "service": "CA",
        "currency": "JPY", "effective_week_start": "2026-05-25",
        "day1": 10, "day2": 11, "day3": 12, "day4": 13,
        "day5": 14, "day6": 15, "day7": 16, "remark": "wk",
    }]
    content, _ = fill_template("air", rows)  # 无 tier_prices → 周表分支
    path = tmp_path / "air_market_price_filled.xlsx"
    path.write_bytes(content)
    return path


def _make_origin_layout_sheet(tmp_path):
    """手造「带起运港列」的老布局周表，用于验证保留的 AirWeeklyAdapter 仍可用。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "May 25 to May 31"
    headers = [
        "Origin (POL)", "Destinations", "Service/+100KG",
        "2026/5/25 (Mon)", "2026/5/26 (Tue)", "2026/5/27 (Wed)",
        "2026/5/28 (Thu)", "2026/5/29 (Fri)", "2026/5/30 (Sat)",
        "2026/5/31 (Sun)", "Remark (Selling)", None, "Currency",
    ]
    for c, h in enumerate(headers, start=1):
        ws.cell(1, c).value = h
    data = ["PVG", "NRT", "CA", 10, 11, 12, 13, 14, 15, 16, "wk", None, "JPY"]
    for c, v in enumerate(data, start=1):
        ws.cell(2, c).value = v
    path = tmp_path / "air_weekly_origin_layout.xlsx"
    wb.save(path)
    return path


def test_air_weekly_priority_before_air(tmp_path):
    assert AirWeeklyAdapter().priority < AirAdapter().priority


def test_make_sheet_roundtrips_via_air_adapter(tmp_path):
    """严格按模板后做表周表无起运港列 → AirWeeklyAdapter 不再命中，回流由 AirAdapter 接管：
    起运港默认补 PVG、币种默认 CNY、目的港/价/周期无损。"""
    path = _make(tmp_path)
    assert AirWeeklyAdapter().detect(path) is False  # 无 origin 表头 → 不认领

    batch = AirAdapter().parse(path, db=None)
    assert batch.file_type is Step1FileType.air
    recs = [r for r in batch.records if r.record_kind == "air_weekly"]
    assert len(recs) == 1
    r = recs[0]
    assert r.origin_port_name == "PVG"   # AirAdapter 默认补
    assert r.destination_port_name == "NRT"
    # 模板表头 Service/+100KG 的档位标记前置进服务描述(2026-06-10 邓老师要求可见)
    assert r.service_desc == "+100KG · CA"
    assert r.currency == "CNY"           # 模板无币种列 → 默认 CNY
    rate = to_air_freight_rate(r, uuid.uuid4())
    assert rate.destination == "NRT"
    assert str(rate.price_day1) == "10"
    assert str(rate.price_day7) == "16"


def test_air_weekly_adapter_still_parses_origin_layout(tmp_path):
    """AirWeeklyAdapter 旧码保留备查：手动上传「带起运港列」的老布局周表仍能识别，
    且币种无损（带 Currency 列）。(默认做表产物已不走这条路。)"""
    path = _make_origin_layout_sheet(tmp_path)
    assert AirWeeklyAdapter().detect(path) is True

    batch = AirWeeklyAdapter().parse(path, db=None)
    recs = [r for r in batch.records if r.record_kind == "air_weekly"]
    assert len(recs) == 1
    r = recs[0]
    assert r.origin_port_name == "PVG"
    assert r.destination_port_name == "NRT"
    assert r.service_desc == "CA"
    assert r.currency == "JPY"           # 带 currency 列 → 币种无损
    assert str(r.effective_week_start) == "2026-05-25"
    rate = to_air_freight_rate(r, uuid.uuid4())
    assert str(rate.price_day1) == "10"
    assert str(rate.price_day7) == "16"
