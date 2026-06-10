from pathlib import Path
from openpyxl import Workbook
from app.services.step1_rates.adapters.air_tier import AirTierAdapter
from app.services.step1_rates.entities import Step1FileType


def _make_tier_xlsx(tmp_path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Air Rates 2026-06-01"
    header = ["Origin (POL)", "Destination", "Service", "45KG", "100KG",
              "Currency", "Effective From", "Effective To",
              "Carrier", "Cargo Class", "Packing", "Density", "Remark"]
    for c, label in enumerate(header, start=1):
        ws.cell(1, c).value = label
    ws.append(["PVG", "NRT", "CA", 17, 14, "JPY",
               "2026-06-01", "2026-06-07", "CA", "普货", "托", "1:167", "周一报价"])
    path = tmp_path / "air_tier_rate_sheet_filled.xlsx"
    wb.save(path)
    return path


def _make_ocean_xlsx(tmp_path: Path) -> Path:
    wb = Workbook()
    wb.active.title = "JP N RATE FCL & LCL"
    path = tmp_path / "ocean_xx.xlsx"
    wb.save(path)
    return path


def test_detect_by_content(tmp_path):
    a = AirTierAdapter()
    assert a.detect(_make_tier_xlsx(tmp_path)) is True
    assert a.detect(_make_ocean_xlsx(tmp_path)) is False


def test_detect_by_hint(tmp_path):
    a = AirTierAdapter()
    assert a.detect(_make_ocean_xlsx(tmp_path), file_type_hint=Step1FileType.air_tier) is True


def test_parse_tier_rows(tmp_path):
    batch = AirTierAdapter().parse(_make_tier_xlsx(tmp_path), db=None)
    assert batch.file_type is Step1FileType.air_tier
    assert len(batch.records) == 1
    r = batch.records[0]
    assert r.record_kind == "air_tier"
    assert r.origin_port_name == "PVG"
    assert r.destination_port_name == "NRT"
    assert r.service_desc == "CA"
    assert r.currency == "JPY"
    assert str(r.valid_from) == "2026-06-01"
    assert r.extras["tier_prices"] == {45: 17.0, 100: 14.0}
    assert r.extras["cargo_class"] == "普货"
    assert r.extras["carrier"] == "CA"


def test_parse_skips_non_numeric_tier(tmp_path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Air Rates"
    for c, label in enumerate(["Origin (POL)", "Destination", "Service", "45KG", "100KG"], start=1):
        ws.cell(1, c).value = label
    ws.append(["PVG", "NRT", "CA", "ASK", 14])   # 45KG 非数字 → 跳过, 不炸
    path = tmp_path / "air_tier_dirty.xlsx"
    wb.save(path)
    batch = AirTierAdapter().parse(path, db=None)
    assert len(batch.records) == 1
    assert batch.records[0].extras["tier_prices"] == {100: 14.0}


def test_parse_slash_date(tmp_path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    for c, label in enumerate(["Origin (POL)", "Destination", "45KG", "Effective From"], start=1):
        ws.cell(1, c).value = label
    ws.append(["PVG", "NRT", 17, "2026/06/01"])
    path = tmp_path / "air_tier_slash.xlsx"
    wb.save(path)
    rec = AirTierAdapter().parse(path, db=None).records[0]
    assert str(rec.valid_from) == "2026-06-01"


def _make_ees_xlsx(tmp_path: Path) -> Path:
    """模拟真实 EES 报价表：封面 sheet + 日本线 sheet(中文档位表头、
    目的港/航司合并单元格留空、'/' 无档、'议价' 非数字价)。"""
    wb = Workbook()
    cover = wb.active
    cover.title = "封面"
    cover["B2"] = "上海百福东方国际物流有限责任公司"
    ws = wb.create_sheet("日本线")
    ws.append(["", "**本运价于2026年5月1日生效**"])
    ws.append(["", "目的港", "航班", "比重", "≧45KG", "≧100KG", "≧500KG", "≧1000KG", "备注"])
    ws.append(["", "KIX", "CK/MU", None, 17, 14, "/", "/", "泡货:3/7分"])
    ws.append(["", None, None, "托盘1:200", "/", "/", 13.5, 13, None])
    ws.append(["", None, "NH中转", None, "/", 16, "/", "/", None])
    ws.append(["", "NRT", "CK/MU", None, 16.5, 13.5, "/", "/", None])
    ws.append(["", None, "FX", None, "议价", "/", 13, "/", None])
    path = tmp_path / "EES（2026-5-21）报价.xlsx"
    wb.save(path)
    return path


def test_detect_ees_chinese_tier_sheet(tmp_path):
    """EES 中文档位表头(目的港/≧100KG)须被识别——此前只认英文契约导致静默 0 入库。"""
    assert AirTierAdapter().detect(_make_ees_xlsx(tmp_path)) is True


def test_parse_ees_chinese_tier_sheet(tmp_path):
    batch = AirTierAdapter().parse(_make_ees_xlsx(tmp_path), db=None)
    assert batch.file_type is Step1FileType.air_tier
    assert len(batch.records) == 5
    assert all(r.record_kind == "air_tier" for r in batch.records)

    r0 = batch.records[0]
    assert r0.destination_port_name == "KIX"
    assert r0.extras["tier_prices"] == {45: 17.0, 100: 14.0}
    assert r0.extras["carrier"] == "CK/MU"
    assert r0.currency == "CNY"
    assert str(r0.valid_from) == "2026-05-21"  # 文件名报价日

    r1 = batch.records[1]  # 目的港/航司合并单元格 → 前向填充
    assert r1.destination_port_name == "KIX"
    assert r1.extras["carrier"] == "CK/MU"
    assert r1.service_desc == "托盘1:200"
    assert r1.extras["tier_prices"] == {500: 13.5, 1000: 13.0}

    r3 = batch.records[3]
    assert r3.destination_port_name == "NRT"
    assert r3.extras["tier_prices"] == {45: 16.5, 100: 13.5}

    r4 = batch.records[4]  # '议价' 不算价，只收 500 档
    assert r4.extras["tier_prices"] == {500: 13.0}


def test_parse_english_contract_still_works_after_ees_support(tmp_path):
    """英文契约表头(做表动态档位表回流)解析行为不受 EES 支持影响。"""
    batch = AirTierAdapter().parse(_make_tier_xlsx(tmp_path), db=None)
    assert len(batch.records) == 1
    assert batch.records[0].extras["tier_prices"] == {45: 17.0, 100: 14.0}


def test_parse_no_match_warns(tmp_path):
    from openpyxl import Workbook
    wb = Workbook()
    wb.active.title = "JP N RATE FCL & LCL"   # 非档位表
    path = tmp_path / "not_tier.xlsx"
    wb.save(path)
    batch = AirTierAdapter().parse(path, db=None)
    assert batch.records == []
    assert batch.warnings  # 落空有提示
