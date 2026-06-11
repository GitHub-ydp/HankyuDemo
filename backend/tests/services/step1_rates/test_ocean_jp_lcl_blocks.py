"""P1-3 回归：JP sheet（FCL sheet）内嵌的 "LCL NORMAL RATE" 区块必须被采集。

2026-06-06 链路A 实测：JP N RATE FCL & LCL 内的国内 LCL 区块共 11 行
（From: Shanghai 10 行 + From: Taicang 1 行 Ferry Shimonoseki）整体漏入，
lcl_rates 只有独立 LCL N RATE sheet 的 28 行。
区块列布局与独立 LCL sheet 不同：To=A, Ocean Freight(CBM/TON)=B,
Surcharges=C, Sailing Day=E, Via=G, Transit Time=I, RMKS=K。
"""
from __future__ import annotations

from datetime import datetime
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook

from app.services.step1_rates.adapters.ocean import OceanAdapter


def _make_jp_like_worksheet():
    wb = Workbook()
    ws = wb.active
    ws.title = "JP N RATE FCL & LCL"
    ws["A3"] = "Effective from"
    ws["B3"] = datetime(2026, 4, 1)
    ws["C3"] = "to"
    ws["D3"] = datetime(2026, 4, 30)
    ws["A7"] = "From: Shanghai"

    # 第一个 LCL 区块（From: Shanghai）
    ws["A114"] = "LCL NORMAL RATE "
    ws["A115"] = "From: Shanghai"
    ws["B115"] = "Currency: USD "
    ws["A116"] = "To"
    ws["B116"] = "Ocean Freight      (CBM/TON)"
    ws["C116"] = "Surcharges"
    ws["E116"] = "Sailig Day"
    ws["G116"] = "Via"
    ws["I116"] = "Transit Time"
    ws["K116"] = "RMKS"
    ws["A117"] = "TOKYO"
    ws["B117"] = 10
    ws["C117"] = "Subject to Destination Charges"
    ws["E117"] = "1,5,7"
    ws["G117"] = "DIRECT"
    ws["I117"] = "4days"
    ws["K117"] = "Mon/Fri HHE own Consol"
    ws["A118"] = "TAKAMATSU"
    ws["B118"] = 120
    ws["C118"] = "Subject to Destination Charges"
    ws["E118"] = "6"
    ws["G118"] = "DIRECT"
    ws["I118"] = "3days"
    ws["K118"] = "MINSHEN"

    # 第二个 LCL 区块（From: Taicang）
    ws["A128"] = "LCL NORMAL RATE "
    ws["A129"] = "From: Taicang"
    ws["B129"] = "Currency: USD "
    ws["A130"] = "To"
    ws["B130"] = "Ocean Freight      (CBM/TON)"
    ws["C130"] = "Surcharges"
    ws["E130"] = "Sailig Day"
    ws["G130"] = "Via"
    ws["I130"] = "Transit Time"
    ws["K130"] = "RMKS"
    ws["A131"] = "Ferry Shimonoseki"
    ws["B131"] = 69
    ws["C131"] = "Subject to Destination Charges"
    ws["E131"] = "MON / FRI"
    ws["G131"] = "DIRECT"
    ws["I131"] = "2days"
    ws["K131"] = "SSF"
    return ws


def test_inline_lcl_blocks_are_collected():
    adapter = OceanAdapter()
    ws = _make_jp_like_worksheet()
    records, warnings = adapter._parse_inline_lcl_blocks(ws, None, source_file="ocean.xlsx")

    assert len(records) == 3
    by_dest = {r.destination_port_name: r for r in records}

    tokyo = by_dest["TOKYO"]
    assert tokyo.record_kind == "lcl"
    assert tokyo.origin_port_name == "Shanghai"
    assert tokyo.freight_per_cbm == Decimal("10")
    assert tokyo.freight_per_ton == Decimal("10")
    assert tokyo.currency == "USD"
    assert tokyo.valid_from is not None and tokyo.valid_from.isoformat() == "2026-04-01"
    assert tokyo.valid_to is not None and tokyo.valid_to.isoformat() == "2026-04-30"
    assert tokyo.sailing_day == "1,5,7"
    assert tokyo.via == "DIRECT"
    assert tokyo.transit_time_text == "4days"
    assert "Mon/Fri HHE own Consol" in (tokyo.remarks or "")
    assert "Subject to Destination Charges" in (tokyo.remarks or "")

    takamatsu = by_dest["TAKAMATSU"]
    assert takamatsu.freight_per_cbm == Decimal("120")

    ferry = by_dest["Ferry Shimonoseki"]
    assert ferry.origin_port_name == "Taicang"
    assert ferry.freight_per_cbm == Decimal("69")


REAL_OCEAN_FILE = (
    Path(__file__).resolve().parents[4]
    / "资料"
    / "2026.04.21"
    / "RE_ 今後の進め方に関するご提案"
    / "【Ocean】 Sea Net Rate_2026_Apr.21 - Apr.30.xlsx"
)


def test_real_ocean_file_collects_11_inline_lcl_rows():
    if not REAL_OCEAN_FILE.exists():
        pytest.skip(f"Ocean 真实样本不可用：{REAL_OCEAN_FILE}")
    batch = OceanAdapter().parse(REAL_OCEAN_FILE)
    jp_lcl = [
        r
        for r in batch.records
        if r.record_kind == "lcl" and r.extras.get("sheet_name") == "JP N RATE FCL & LCL"
    ]
    assert len(jp_lcl) == 11
    dests = [r.destination_port_name for r in jp_lcl]
    assert "TOKYO" in dests
    assert "Ferry Shimonoseki" in dests
    taicang = [r for r in jp_lcl if (r.origin_port_name or "").startswith("Taicang")]
    assert len(taicang) == 1
    # 独立 LCL sheet 原有 28 行不受影响
    standalone = [
        r
        for r in batch.records
        if r.record_kind == "lcl" and r.extras.get("sheet_name") == "LCL N RATE"
    ]
    assert len(standalone) == 28
