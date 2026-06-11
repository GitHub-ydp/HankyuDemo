"""指定数据下载产物回流导入：文件名不含 ocean 也应按 sheet 内容识别并解析。

场景：做表页「指定数据下载」把会话 rows 回填进用户上传的模板，产物命名
`<模板名>_filled.xlsx`（不含 "ocean"）。用户整理后从运价导入页上传时：
- 解析器留「自动」→ 旧逻辑只认文件名 → LookupError → 错走 AI 兜底报错；
- 用户整理时可能删掉 JP/LCL sheet → 旧 parse 硬取三表 → KeyError。
"""
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from app.services.step1_rates.adapters import OceanAdapter
from app.services.step1_rates.entities import Step1FileType
from app.services.step1_rates.service import build_default_registry
from app.services.step1_rates.sheet_builder.template_refill import refill_into_template

FIXTURE = (
    Path(__file__).resolve().parents[2] / "fixtures" / "sea_other_ports_template.xlsx"
)

ROWS = [
    {"destination": "BUSAN", "carrier": "EAS", "container_20gp": 160, "container_40hq": 320},
    {"destination": "HONG KONG", "carrier": "SITC", "container_20gp": 100, "container_40hq": 200},
]


def _refilled_file(tmp_path: Path, template_bytes: bytes, name: str) -> Path:
    out = tmp_path / name
    out.write_bytes(refill_into_template(template_bytes, ROWS))
    return out


def test_registry_resolves_refilled_file_without_ocean_in_name(tmp_path):
    """文件名不含 ocean（指定下载产物默认命名）也应命中 OceanAdapter。"""
    path = _refilled_file(tmp_path, FIXTURE.read_bytes(), "sea_other_ports_template_filled.xlsx")
    registry = build_default_registry()
    adapter = registry.resolve(path)
    assert adapter.key == "ocean"


def test_parse_refilled_file_yields_rows(tmp_path):
    """指定下载产物（含三表）无 hint 直接解析出运价行。"""
    path = _refilled_file(tmp_path, FIXTURE.read_bytes(), "整理后的运价_filled.xlsx")
    batch = build_default_registry().parse(path, db=None)
    assert batch.adapter_key == "ocean"
    assert batch.records
    destinations = {r.destination_port_name for r in batch.records}
    assert "BUSAN" in destinations


def test_parse_tolerates_missing_jp_and_lcl_sheets(tmp_path):
    """用户模板/整理后只剩 OTHER PORTS 一张表：有哪张解析哪张，不许 KeyError。"""
    wb = load_workbook(BytesIO(FIXTURE.read_bytes()))
    del wb["JP N RATE FCL & LCL"]
    del wb["LCL N RATE"]
    buf = BytesIO()
    wb.save(buf)
    path = _refilled_file(tmp_path, buf.getvalue(), "only_other_ports_filled.xlsx")

    batch = build_default_registry().parse(path, db=None)
    assert batch.adapter_key == "ocean"
    assert batch.records
    assert all(r.extras.get("sheet_name") == "FCL N RATE OF OTHER PORTS" for r in batch.records)


def test_detect_does_not_claim_unrelated_xlsx(tmp_path):
    """无海运特征 sheet 的 xlsx 不应被内容识别误吞。"""
    wb = Workbook()
    wb.active.title = "Sheet1"
    path = tmp_path / "random_filled.xlsx"
    wb.save(path)
    assert OceanAdapter().detect(path) is False


def test_detect_survives_non_xlsx_file(tmp_path):
    """非 xlsx（如 .msg/.pdf）走内容识别不许抛异常，返回 False 即可。"""
    path = tmp_path / "mail.msg"
    path.write_bytes(b"not an excel")
    assert OceanAdapter().detect(path) is False


def test_hinted_parse_without_any_ocean_sheet_raises_friendly_error(tmp_path):
    """手选海运但文件缺全部海运表：给可读 ValueError（路由转 400），不许 KeyError。"""
    wb = Workbook()
    wb.active.title = "Sheet1"
    path = tmp_path / "whatever.xlsx"
    wb.save(path)
    with pytest.raises(ValueError, match="海运"):
        OceanAdapter().parse(path, db=None)
