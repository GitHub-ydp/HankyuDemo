"""P0-1: 适配器命中但解析 0 行时，上传必须显式报错（NoRatesFoundError → API 422），
禁止静默建空 draft（用户上传→以为成功→activate=empty_batch→数据全丢无感知）。"""
from pathlib import Path

import pytest
from openpyxl import Workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker

from app.models.base import Base
from app.services import rate_batch_service


@pytest.fixture()
def db() -> Session:
    engine = create_engine(
        "sqlite:///:memory:", connect_args={"check_same_thread": False}
    )
    Base.metadata.create_all(engine)
    s = sessionmaker(bind=engine, autoflush=False, autocommit=False)()
    try:
        yield s
    finally:
        s.close()


def _header_only_tier_xlsx() -> bytes:
    """英文契约表头但无任何数据行：AirTierAdapter detect=True、parse=0 行。"""
    import io

    wb = Workbook()
    ws = wb.active
    ws.title = "Air Rates"
    for c, label in enumerate(
        ["Origin (POL)", "Destination", "Service", "45KG", "100KG"], start=1
    ):
        ws.cell(1, c).value = label
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_upload_zero_rows_raises_no_rates_found(db, tmp_path, monkeypatch):
    monkeypatch.setattr(
        rate_batch_service.settings, "upload_dir", str(tmp_path)
    )
    with pytest.raises(rate_batch_service.NoRatesFoundError):
        rate_batch_service.create_draft_batch_from_upload(
            file_name="air_tier_header_only.xlsx",
            content=_header_only_tier_xlsx(),
            db=db,
        )
    # 报错路径不留孤儿文件
    assert list(Path(tmp_path).glob("step1_batch_*")) == []


def _ees_xlsx_bytes() -> bytes:
    """EES 风格中文档位表（目的港/航班/比重/≧45…KG），见 FIXLIST P0-2。"""
    import io

    wb = Workbook()
    wb.active.title = "封面"
    ws = wb.create_sheet("日本线")
    ws.append(["", "目的港", "航班", "比重", "≧45KG", "≧100KG", "≧500KG", "≧1000KG"])
    ws.append(["", "KIX", "CK/MU", None, 17, 14, "/", "/"])
    ws.append(["", None, None, "托盘1:200", "/", "/", 13.5, 13])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_ees_upload_to_air_tier_rates_end_to_end(db, tmp_path, monkeypatch):
    """P0-2 验收：导入 EES 中文档位表 → 激活 → air_tier_rates > 0（此前静默 0 入库）。"""
    from app.models.air_tier_rate import AirTierRate

    monkeypatch.setattr(
        rate_batch_service.settings, "upload_dir", str(tmp_path)
    )
    detail = rate_batch_service.create_draft_batch_from_upload(
        file_name="EES（2026-5-21）报价.xlsx",
        content=_ees_xlsx_bytes(),
        db=db,
    )
    assert detail["total_rows"] == 2

    result = rate_batch_service.activate_rate_batch(
        detail["batch_id"], db, dry_run=False
    )
    assert result["activation_status"] == "activated"

    rows = db.query(AirTierRate).all()
    assert len(rows) == 2
    first = next(r for r in rows if r.tier_prices.get(45) or r.tier_prices.get("45"))
    assert first.destination == "KIX"
    assert first.origin == "PVG"  # EES 上海起运默认 PVG
    assert first.currency == "CNY"
    assert first.carrier == "CK/MU"
