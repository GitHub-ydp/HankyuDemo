import zipfile
from decimal import Decimal
from pathlib import Path

import extract_msg
import pytest
from openpyxl import load_workbook

from app.services.step2_bidding.customer_profiles.nitori import NitoriProfile
from app.services.step2_bidding.entities import RowStatus
from app.services.step2_bidding.nitori_cost_book import NitoriCostBook

NITORI_DIR = Path(__file__).resolve().parents[4] / "资料" / "2026.05.26" / "_nitori_unzip" / "ニトリ様海上入札"
QUOTE_GLOBAL = NITORI_DIR / "_【to GLOBAL】2026年7月～9月_見積り書.xlsm"
COST_MSG = NITORI_DIR / "②（Cost）回复 HHESHA内部　転送 ニトリ海上運賃入札（2026年2Q7－９月）.msg"


def _cost_book(tmp_path):
    m = extract_msg.Message(str(COST_MSG))
    for a in m.attachments:
        n = a.longFilename or a.shortFilename or ""
        if n.lower().endswith(".xlsx"):
            p = tmp_path / n
            p.write_bytes(a.data)
            m.close()
            return NitoriCostBook.from_xlsx(p)
    raise AssertionError("no cost xlsx")


def test_nitori_detect():
    assert NitoriProfile().detect(QUOTE_GLOBAL) is True


def test_nitori_parse_china_rows():
    parsed = NitoriProfile().parse(QUOTE_GLOBAL, bid_id="b1", period="2026Q2")
    china = [r for r in parsed.rows if r.origin_code in ("SHANGHAI", "TAICANG")]
    assert len(china) == 48
    r114 = next(r for r in parsed.rows if r.row_idx == 114)
    assert r114.origin_code == "SHANGHAI"
    assert r114.destination_code == "PORT KLANG"
    assert r114.extras["size"] == "20F"
    assert r114.extras["pod_raw"] == "PORT KLANG (NORTH)"
    assert r114.extras["is_china"] is True


def test_nitori_match_filled_and_norate(tmp_path):
    prof = NitoriProfile(cost_book=_cost_book(tmp_path))
    parsed = prof.parse(QUOTE_GLOBAL, bid_id="b1", period="2026Q2")
    reports = prof.match(parsed)
    by_row = {rp.row_idx: rp for rp in reports}
    # r114 SHANGHAI->PORT KLANG(NORTH) 20F 命中 580
    assert by_row[114].status == RowStatus.FILLED
    assert by_row[114].cost_price == Decimal("580")
    assert by_row[114].sell_price == Decimal("667")          # 580*1.15=667
    # r120 TAICANG->PORT KLANG 成本 NO SERVICE → no_rate
    assert by_row[120].status == RowStatus.NO_RATE
    # 非中国发行不进 report（如第 9 行 OSAKA）
    assert 9 not in by_row


def test_nitori_fill_cost_and_sr(tmp_path):
    prof = NitoriProfile(cost_book=_cost_book(tmp_path))
    parsed = prof.parse(QUOTE_GLOBAL, bid_id="b1", period="2026Q2")
    reports = prof.match(parsed)
    cost_out = tmp_path / "cost.xlsm"
    sr_out = tmp_path / "sr.xlsm"
    prof.fill(QUOTE_GLOBAL, parsed, reports, "cost", cost_out)
    prof.fill(QUOTE_GLOBAL, parsed, reports, "sr", sr_out)

    wbc = load_workbook(cost_out, keep_vba=True)["Quotation (Global) Jul-Sep"]
    wbs = load_workbook(sr_out, keep_vba=True)["Quotation (Global) Jul-Sep"]
    assert wbc.cell(114, 19).value == 580        # 成本版 OCEAN FREIGHT amount
    assert wbs.cell(114, 19).value == 667        # 报价版 = 580*1.15
    assert wbc.cell(114, 18).value == "USD"      # cur
    # no_rate 行(r120)不写运价
    assert wbc.cell(120, 19).value in (None, 0, "")
    # 宏保留
    assert any("vbaProject" in n for n in zipfile.ZipFile(cost_out).namelist())


def test_nitori_fill_rejects_bad_variant(tmp_path):
    prof = NitoriProfile(cost_book=_cost_book(tmp_path))
    parsed = prof.parse(QUOTE_GLOBAL, bid_id="b1", period="2026Q2")
    with pytest.raises(ValueError):
        prof.fill(QUOTE_GLOBAL, parsed, [], "bad", tmp_path / "x.xlsm")


def test_nitori_fill_writes_carrier_thc_doc_lss(tmp_path):
    """fill 回写 CARRIER/THC/DOC/LSS 列；LSS 未知不再硬编 USD 0。"""
    from app.services.step2_bidding.entities import PerRowReport

    prof = NitoriProfile()
    parsed = prof.parse(QUOTE_GLOBAL, bid_id="b1", period="2026Q2")

    def _rep(row_idx, **kw):
        base = dict(
            row_idx=row_idx, section_code="GLOBAL", destination_code="X",
            status=RowStatus.FILLED, cost_price=Decimal("300"),
            sell_price=Decimal("345"), markup_ratio=Decimal("1.15"),
            lead_time_text="2 DAYS", carrier_text="COSCO", remark_text=None,
            selected_candidate=None,
        )
        base.update(kw)
        return PerRowReport(**base)

    reports = [
        _rep(114, thc_amount=Decimal("982"), doc_amount=Decimal("450"),
             lss_amount=Decimal("200")),
        _rep(115, carrier_text="EMC"),     # 附加费全未知
    ]
    out = tmp_path / "o.xlsm"
    prof.fill(QUOTE_GLOBAL, parsed, reports, "cost", out)
    ws = load_workbook(out, keep_vba=True)["Quotation (Global) Jul-Sep"]
    assert ws.cell(114, 4).value == "COSCO"            # CARRIER
    assert ws.cell(114, 26).value == "CNY"             # THC cur
    assert ws.cell(114, 27).value == 982               # THC amount
    assert ws.cell(114, 28).value == "CNY"             # DOC cur
    assert ws.cell(114, 29).value == 450               # DOC amount
    assert ws.cell(114, 20).value == "USD"             # LSS cur
    assert ws.cell(114, 21).value == 200               # LSS amount
    assert ws.cell(115, 4).value == "EMC"
    assert ws.cell(115, 21).value is None              # LSS 未知 → 不写 0
    assert ws.cell(115, 20).value is None
    assert ws.cell(115, 27).value is None              # THC 未知 → 不写
