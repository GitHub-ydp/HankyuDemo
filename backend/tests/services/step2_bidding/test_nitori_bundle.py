import shutil
from pathlib import Path
from openpyxl import load_workbook
from app.services.step2_bidding.nitori_bundle import resolve_bundle

NITORI_DIR = Path(__file__).resolve().parents[4] / "资料" / "2026.05.26" / "_nitori_unzip" / "ニトリ様海上入札"


def test_resolve_bundle(tmp_path):
    work = tmp_path / "pkg"
    shutil.copytree(NITORI_DIR, work)
    quote, cost = resolve_bundle(work)
    assert "GLOBAL" in quote.name
    assert quote.suffix == ".xlsm"
    assert cost.suffix == ".xlsx" and cost.exists()
    # 关键：选中的必须是真正的成本表（含 FCL sheet），不是 ① 邮件里的旧报价/合同
    wb = load_workbook(cost, read_only=True)
    try:
        assert "FCL" in wb.sheetnames
    finally:
        wb.close()
