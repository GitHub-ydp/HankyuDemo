from io import BytesIO
from openpyxl import load_workbook
from app.services.step1_rates.sheet_builder.template_filler import fill_template


def test_sea_sheet_splits_40gp_40hq_and_meta():
    rows = [{
        "destination": "USLAX", "carrier": "ONE",
        "container_20gp": 1000, "container_40gp": 1800, "container_40hq": 1850,
        "freight_20": 1000, "freight_40": 1800,
        "currency": "USD", "valid_from": "2026-06-01", "valid_to": "2026-06-30",
        "rate_level": "NAC", "service_code": "EC1", "remark": "test",
    }]
    content, _ = fill_template("sea", rows)
    ws = load_workbook(BytesIO(content))["JP N RATE FCL & LCL"]
    # 第 8 行表头含新列
    header = [c.value for c in ws[8]]
    for label in ["Currency", "Valid From", "Valid To", "Rate Level", "Service Code"]:
        assert label in header, f"缺表头 {label}"
    # 数据从第 9 行起，3 行：20FT/40GP/40HQ，箱型价各异
    labels = [ws.cell(9 + i, 3).value for i in range(3)]
    assert labels == ["20FT", "40GP", "40HQ"]
    freights = [ws.cell(9 + i, 4).value for i in range(3)]
    assert freights == [1000, 1800, 1850]   # 40GP≠40HQ 不再合并
    # 元数据每行都带（取第 9 行）
    cur_col = header.index("Currency") + 1
    assert ws.cell(9, cur_col).value == "USD"
