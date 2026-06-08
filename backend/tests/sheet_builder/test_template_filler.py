"""模板填充器测试：把 normalized rate dict 填进空白模板的正确格子，且不破坏表头。"""
from io import BytesIO

from openpyxl import load_workbook

from app.services.step1_rates.sheet_builder.template_filler import fill_template


def _reload(content: bytes):
    return load_workbook(BytesIO(content), data_only=False)


def test_fill_sea_expands_container_rows():
    rows = [
        {
            "destination": "BUSAN",
            "carrier": "SJJ",
            "container_20gp": 130,
            "container_40gp": 260,
            "container_40hq": 265,
            "lss_cic": "Incl.",
            "baf": 50,
            "transit": "2days",
            "remark": "直达",
        },
        {"destination": "INCHON", "carrier": "COSCO",
         "container_20gp": 300, "container_40gp": 500, "container_40hq": 520},
    ]
    content, filename = fill_template("sea", rows)
    ws = _reload(content)["JP N RATE FCL & LCL"]

    # 表头未被破坏
    assert ws.cell(8, 1).value == "To"
    # 第一条 BUSAN 展开 20FT / 40GP / 40HQ 三行，从 r9 开始
    assert ws.cell(9, 1).value == "BUSAN"
    assert ws.cell(9, 2).value == "SJJ"
    assert ws.cell(9, 3).value == "20FT"
    assert ws.cell(9, 4).value == 130
    assert ws.cell(9, 5).value == "Incl."  # LSS+CIC
    assert ws.cell(9, 6).value == 50       # BAF
    assert ws.cell(10, 3).value == "40GP"
    assert ws.cell(10, 4).value == 260
    assert ws.cell(11, 3).value == "40HQ"
    assert ws.cell(11, 4).value == 265
    # 第二条 INCHON 从 r12，同样三行
    assert ws.cell(12, 1).value == "INCHON"
    assert ws.cell(12, 3).value == "20FT"
    assert ws.cell(12, 4).value == 300
    assert ws.cell(13, 4).value == 500   # 40GP
    assert ws.cell(14, 4).value == 520   # 40HQ
    assert filename.endswith(".xlsx")


def test_fill_air_one_row_per_rate():
    rows = [
        {
            "origin": "PVG",
            "destination": "NRT",
            "service": "普货",
            "day1": 14,
            "day2": 14,
            "day7": 15,
            "remark": "x",
        }
    ]
    content, _ = fill_template("air", rows)
    ws = _reload(content)["May 25 to May 31"]

    # 严格按客户原件布局(无起运港列)：A 目的港 / B 服务 / C-I 每日价 / J 备注
    assert ws.cell(1, 1).value == "Destinations"  # 表头未破坏，A 列即目的港
    assert ws.cell(1, 2).value == "Service/+100KG"
    assert ws.cell(2, 1).value == "NRT"
    assert ws.cell(2, 2).value == "普货"
    assert ws.cell(2, 3).value == 14   # day1（C 列）
    assert ws.cell(2, 9).value == 15   # day7（I 列）
    assert ws.cell(2, 10).value == "x"  # 备注（J 列）
    assert "Origin" not in [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]


def test_fill_empty_rows_keeps_header_and_blank_body():
    content, _ = fill_template("sea", [])
    ws = _reload(content)["JP N RATE FCL & LCL"]
    assert ws.cell(8, 1).value == "To"            # 表头还在
    assert ws.cell(9, 1).value in (None, "")      # 数据区仍空


def test_fill_air_rewrites_week_headers_from_data():
    """行带 effective_week_start 时，按该周改写 day1-7 日期表头 + sheet 名。"""
    rows = [
        {"destination": "NRT", "service": "CK", "day1": 14, "day7": 14,
         "effective_week_start": "2026-06-01"},
    ]
    content, _ = fill_template("air", rows)
    wb = _reload(content)

    assert wb.sheetnames == ["Jun 1 to Jun 7"], "sheet 名应按周改写"
    ws = wb["Jun 1 to Jun 7"]
    assert str(ws.cell(1, 3).value).startswith("2026/6/1"), "day1 表头应为该周起始日(C 列)"
    assert str(ws.cell(1, 9).value).startswith("2026/6/7"), "day7 表头应为该周第7日(I 列)"
    assert ws.cell(2, 1).value == "NRT"  # 数据照填(目的港在 A 列)


def test_fill_air_without_week_keeps_template():
    """无 effective_week_start（如重量档报价）→ 保留模板原日期/名，不乱改。"""
    rows = [{"destination": "NRT", "service": "CK", "day1": 14}]
    content, _ = fill_template("air", rows)
    wb = _reload(content)

    assert wb.sheetnames == ["May 25 to May 31"], "无周信息应保留模板原样"
    assert str(wb["May 25 to May 31"].cell(1, 3).value).startswith("2026/5/25")  # day1 表头 C 列


def test_build_air_tier_sheet_dynamic_columns():
    """档位行(带 tier_prices)显式关掉 strict → 程序生成动态档位表(不套 air_blank)：
    起运港|目的港|服务|动态 KG 列(全表并集升序)|备注；稀疏档某行缺的留空。
    (默认已改为严格按模板，旧多档行为保留于 strict_air_template=False。)"""
    rows = [
        {"origin": "PVG", "destination": "KIX", "service": "CK/MU",
         "tier_prices": {45: 17.0, 100: 14.0}, "remark": "含油"},
        {"origin": "PVG", "destination": "BKK", "service": "平散货",
         "tier_prices": {100: 16.0, 300: 16.0, 500: 16.0, 1000: 16.0}, "remark": None},
    ]
    content, filename = fill_template("air", rows, strict_air_template=False)
    wb = _reload(content)
    ws = wb[wb.sheetnames[0]]

    # 表头：固定列 + 全表档位并集(45/100/300/500/1000)升序 + 元数据列 + 备注
    header = [ws.cell(1, c).value for c in range(1, 17)]
    assert header == [
        "Origin (POL)", "Destination", "Service",
        "45KG", "100KG", "300KG", "500KG", "1000KG",
        "Currency", "Effective From", "Effective To",
        "Carrier", "Cargo Class", "Packing", "Density", "Remark",
    ]
    # KIX 行(r2)：45=17 / 100=14，其余档留空
    assert ws.cell(2, 1).value == "PVG"
    assert ws.cell(2, 2).value == "KIX"
    assert ws.cell(2, 3).value == "CK/MU"
    assert ws.cell(2, 4).value == 17.0           # 45KG
    assert ws.cell(2, 5).value == 14.0           # 100KG
    assert ws.cell(2, 6).value in (None, "")     # 300KG 留空
    assert ws.cell(2, 16).value == "含油"          # 备注(移到元数据列之后)
    # BKK 行(r3)：100/300/500/1000=16，45 留空
    assert ws.cell(3, 4).value in (None, "")     # 45KG 留空
    assert ws.cell(3, 5).value == 16.0
    assert ws.cell(3, 8).value == 16.0           # 1000KG
    assert filename.endswith(".xlsx")


def test_build_air_tier_sheet_names_from_report_date():
    """报价日(effective_week_start)进 sheet 名；无则用通用名。"""
    rows = [{"origin": "PVG", "destination": "KIX", "service": "CK",
             "tier_prices": {100: 14.0}, "effective_week_start": "2026-05-21"}]
    content, _ = fill_template("air", rows, strict_air_template=False)
    wb = _reload(content)
    assert "2026-05-21" in wb.sheetnames[0]


def test_build_air_tier_sheet_accepts_string_tier_keys():
    """下载 POST 经 JSON 往返后档位键会变字符串('45') → 仍能正确成列取值。"""
    rows = [{"origin": "PVG", "destination": "KIX", "service": "CK",
             "tier_prices": {"45": 17.0, "100": 14.0}, "remark": None}]
    content, _ = fill_template("air", rows, strict_air_template=False)
    ws = _reload(content)[_reload(content).sheetnames[0]]
    header = [ws.cell(1, c).value for c in range(1, 6)]
    assert header[:3] == ["Origin (POL)", "Destination", "Service"]
    assert "45KG" in header and "100KG" in header
    assert ws.cell(2, 4).value == 17.0


def test_strict_air_template_tier_takes_p100_into_weekly():
    """默认严格按模板：档位行降为 air_blank.xlsx 周表行——只取 +100KG 一档，
    铺满 day1-7；套固定模板(sheet 名/列结构与周表一致)，不再是动态档位表。"""
    rows = [{"origin": "PVG", "destination": "KIX", "service": "CK/MU",
             "tier_prices": {45: 17.0, 100: 14.0, 300: 13.0}, "remark": "含油"}]
    content, filename = fill_template("air", rows)  # 默认 strict
    wb = _reload(content)
    ws = wb[wb.sheetnames[0]]
    # 套客户原件周表模板：A 目的港 / B 服务 / C-I 日期；无起运港列
    assert ws.cell(1, 1).value == "Destinations"
    assert ws.cell(1, 2).value == "Service/+100KG"
    assert str(ws.cell(1, 3).value).startswith("2026/5/25")  # 模板默认周表头(C 列, 无周信息)
    # 数据行：目的港/服务到位，价 = 100KG 档(14.0) 铺满 day1-7，45/300 档舍弃
    assert ws.cell(2, 1).value == "KIX"
    assert ws.cell(2, 2).value == "CK/MU"
    for col in range(3, 10):  # day1..day7（C-I）
        assert ws.cell(2, col).value == 14.0
    assert ws.cell(2, 10).value == "含油"  # 备注（J 列）
    assert filename == "air_rate_sheet_filled.xlsx"


def test_strict_air_template_falls_back_to_nearest_tier_when_no_p100():
    """无精确 100KG 档时取最接近 100 的档(45)，并在备注标注实际档位。"""
    rows = [{"origin": "PVG", "destination": "KIX", "service": "CK",
             "tier_prices": {45: 17.0, 300: 13.0}, "remark": "原注"}]
    content, _ = fill_template("air", rows)  # 默认 strict
    ws = _reload(content).active
    assert ws.cell(2, 3).value == 17.0  # day1（C 列）= 45KG 档(距 100 最近)
    assert "45KG" in str(ws.cell(2, 10).value)  # 备注列(J)标注实际档位
    assert "原注" in str(ws.cell(2, 10).value)


def test_strict_air_template_passes_through_weekly_rows():
    """同批里已是周表行(无 tier_prices)原样透传，不被降档逻辑破坏。"""
    rows = [
        {"origin": "PVG", "destination": "NRT", "service": "CK",
         "tier_prices": {100: 20.0}},
        {"origin": "PVG", "destination": "KIX", "service": "MU",
         "day1": 11, "day7": 12},
    ]
    content, _ = fill_template("air", rows)  # 默认 strict
    ws = _reload(content).active
    assert ws.cell(2, 3).value == 20.0   # tier 行 → 100KG 铺满（day1=C 列）
    assert ws.cell(2, 9).value == 20.0   # day7=I 列
    assert ws.cell(3, 3).value == 11     # 周表行原样
    assert ws.cell(3, 9).value == 12
