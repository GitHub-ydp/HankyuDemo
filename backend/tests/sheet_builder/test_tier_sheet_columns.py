from io import BytesIO
from openpyxl import load_workbook
from app.services.step1_rates.sheet_builder.template_filler import fill_template


def test_tier_sheet_has_metadata_columns():
    rows = [{
        "origin": "PVG", "destination": "NRT", "service": "CA",
        "currency": "JPY", "carrier": "CA", "cargo_class": "普货",
        "packing": "托", "density": "1:167",
        "effective_week_start": "2026-06-01", "effective_to": "2026-06-07",
        "remark": "周一报价", "tier_prices": {45: 17.0, 100: 14.0},
    }]
    content, _ = fill_template("air", rows)
    ws = load_workbook(BytesIO(content)).active
    header = [c.value for c in ws[1]]
    for label in ["Origin (POL)", "Destination", "Service", "45KG", "100KG",
                  "Currency", "Effective From", "Effective To",
                  "Carrier", "Cargo Class", "Packing", "Density", "Remark"]:
        assert label in header, f"缺表头列 {label}"
    # 值落对位置
    data = {header[i]: ws.cell(2, i + 1).value for i in range(len(header))}
    assert data["Currency"] == "JPY"
    assert data["Effective From"] == "2026-06-01"
    assert data["Carrier"] == "CA"
    assert data["45KG"] == 17.0
