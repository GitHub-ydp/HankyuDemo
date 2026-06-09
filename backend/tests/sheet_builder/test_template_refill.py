"""按客户模板回填海运运价（OTHER PORTS 页）核心测试。

纯函数 refill_into_template：给模板 bytes + 当前会话 rows，返回填好的 xlsx bytes。
夹具 = 客户真实「空白 着地あり」模板（保留目的港名、价格留空）。
"""
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.services.step1_rates.sheet_builder.template_refill import (
    RefillError,
    refill_into_template,
)

FIXTURE = Path(__file__).resolve().parents[1] / "fixtures" / "sea_other_ports_template.xlsx"
SHEET = "FCL N RATE OF OTHER PORTS"


def _tpl_bytes() -> bytes:
    return FIXTURE.read_bytes()


def _refill(rows):
    content = refill_into_template(_tpl_bytes(), rows)
    assert content[:2] == b"PK"  # xlsx = zip
    return load_workbook(BytesIO(content), data_only=False)


def test_basic_fill_carrier_and_container_rows():
    rows = [
        {"destination": "HONG KONG", "carrier": "SJJ",
         "container_20gp": 230, "container_40gp": 460, "container_40hq": 460,
         "via": "DIRECT", "transit": "4days"},
        {"destination": "HONG KONG", "carrier": "ASL",
         "container_20gp": 240, "container_40hq": 480},
        {"destination": "BUSAN", "carrier": "EAS",
         "container_20gp": 160, "container_40hq": 320},
    ]
    ws = _refill(rows)[SHEET]
    # HONG KONG 从数据起始行 r9 开始，2 船司 × 2 箱型 = 4 行
    assert ws.cell(9, 1).value == "HONG KONG"
    assert ws.cell(9, 2).value == "SJJ"
    assert ws.cell(9, 3).value == "20FT"
    assert ws.cell(9, 4).value == 230
    assert ws.cell(10, 3).value == "40FT/40HQ"
    assert ws.cell(10, 4).value == 460
    assert ws.cell(11, 2).value == "ASL"
    assert ws.cell(11, 4).value == 240
    assert ws.cell(12, 4).value == 480
    # A 列港名竖向合并整块 r9:r12
    assert any(str(m) == "A9:A12" for m in ws.merged_cells.ranges)
    # BUSAN 紧接 r13，1 船司 × 2 行
    assert ws.cell(13, 1).value == "BUSAN"
    assert ws.cell(13, 2).value == "EAS"
    assert ws.cell(14, 3).value == "40FT/40HQ"
    assert ws.cell(14, 4).value == 320


def test_40_row_takes_40hq_when_differ():
    rows = [{"destination": "BUSAN", "carrier": "EAS",
             "container_20gp": 160, "container_40gp": 300, "container_40hq": 320}]
    ws = _refill(rows)[SHEET]
    assert ws.cell(9, 4).value == 160      # 20FT
    assert ws.cell(10, 4).value == 320     # 40 行取 40HQ（≠40GP 时）


def test_40_row_falls_back_to_40gp_when_no_40hq():
    rows = [{"destination": "BUSAN", "carrier": "EAS",
             "container_20gp": 160, "container_40gp": 300}]
    ws = _refill(rows)[SHEET]
    assert ws.cell(10, 4).value == 300     # 40HQ 空 → 回退 40GP


def test_surcharges_map_to_lss_baf_cic_caf():
    rows = [{"destination": "BUSAN", "carrier": "EAS",
             "container_20gp": 160, "container_40hq": 320,
             "surcharges": [
                 {"code": "LSS", "included": True},
                 {"code": "BAF", "payment": "collect"},
                 {"code": "CIC", "amount_20": 50, "amount_40": 100},
                 {"code": "CAF", "note": "subject to dest"},
             ]}]
    ws = _refill(rows)[SHEET]
    # 20FT 行：E=LSS F=BAF G=CIC H=CAF
    assert ws.cell(9, 5).value == "Incl."
    assert ws.cell(9, 6).value == "Collect"
    assert ws.cell(9, 7).value == 50
    assert ws.cell(9, 8).value == "subject to dest"
    # 40 行 CIC 取 amount_40
    assert ws.cell(10, 7).value == 100


def test_surcharge_falls_back_to_flat_lss_cic():
    rows = [{"destination": "BUSAN", "carrier": "EAS", "container_20gp": 160,
             "lss_cic": "Incl.", "baf": 30}]
    ws = _refill(rows)[SHEET]
    assert ws.cell(9, 5).value == "Incl."   # 无 surcharges → 用扁平 lss_cic 落 LSS 列
    assert ws.cell(9, 6).value == 30        # 扁平 baf 落 BAF 列


def test_empty_port_keeps_name_one_blank_row():
    # SINGAPORE 在模板里，rows 没有它 → 留港名 + 1 空行
    rows = [{"destination": "BUSAN", "carrier": "EAS", "container_20gp": 160}]
    ws = _refill(rows)[SHEET]
    col_a = [ws.cell(r, 1).value for r in range(9, 40)]
    assert "SINGAPORE" in col_a
    # SINGAPORE 那一行价格列(D)为空
    sg_row = 9 + col_a.index("SINGAPORE")
    assert ws.cell(sg_row, 4).value in (None, "")


def test_alias_and_multiname_matching():
    rows = [
        {"destination": "PUSAN", "carrier": "EAS", "container_20gp": 160},       # PUSAN↔BUSAN
        {"destination": "CHENNAI", "carrier": "X", "container_20gp": 700},       # ↔ MADRAS / CHENNAI
        {"destination": "LOS ANGELES", "carrier": "Y", "container_20gp": 800},   # ↔ LONG BEACH⏎LOS ANGELES
        {"destination": "CHICAGO", "carrier": "Z", "container_20gp": 900},       # ↔ CHICAGO (via LAX)
    ]
    ws = _refill(rows)[SHEET]
    cells = {(ws.cell(r, 1).value, ws.cell(r, 2).value): ws.cell(r, 4).value
             for r in range(9, ws.max_row + 1)}
    assert ("BUSAN", "EAS") in cells and cells[("BUSAN", "EAS")] == 160
    assert ("MADRAS / CHENNAI", "X") in cells and cells[("MADRAS / CHENNAI", "X")] == 700
    assert ("LONG BEACH\nLOS ANGELES", "Y") in cells and cells[("LONG BEACH\nLOS ANGELES", "Y")] == 800
    assert ("CHICAGO (via LAX)", "Z") in cells and cells[("CHICAGO (via LAX)", "Z")] == 900


def test_unlisted_port_is_filtered_out():
    rows = [{"destination": "DALIAN", "carrier": "EAS", "container_20gp": 999},
            {"destination": "BUSAN", "carrier": "EAS", "container_20gp": 160}]
    ws = _refill(rows)[SHEET]
    col_a = [ws.cell(r, 1).value for r in range(9, ws.max_row + 1)]
    assert "DALIAN" not in col_a       # 模板没列 → 不输出


def test_fidelity_other_sheets_and_header_untouched():
    rows = [{"destination": "BUSAN", "carrier": "EAS", "container_20gp": 160}]
    wb = _refill(rows)
    assert wb.sheetnames == [
        "JP N RATE FCL & LCL", "FCL N RATE OF OTHER PORTS", "LCL N RATE",
    ]
    # 表头 8 行不变
    assert wb[SHEET].cell(8, 1).value == "TO"
    # JP 页原样保留（她模板里 JP 页留了 TOKYO/OSAKA）
    assert wb["JP N RATE FCL & LCL"].cell(9, 1).value == "TOKYO\nYOKOHAMA"


def test_empty_rows_produce_blank_ports_no_crash():
    wb = _refill([])
    ws = wb[SHEET]
    # 所有港留名、价格空
    assert ws.cell(9, 1).value == "HONG KONG"
    assert ws.cell(9, 4).value in (None, "")


def test_missing_sheet_raises_refill_error():
    from openpyxl import Workbook
    buf = BytesIO()
    Workbook().save(buf)            # 全新空 workbook，无 OTHER PORTS 页
    with pytest.raises(RefillError):
        refill_into_template(buf.getvalue(), [])
