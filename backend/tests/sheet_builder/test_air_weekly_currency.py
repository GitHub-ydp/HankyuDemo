from io import BytesIO
from openpyxl import load_workbook
from app.services.step1_rates.sheet_builder.template_filler import fill_template


def test_air_weekly_sheet_omits_origin_and_currency_columns():
    """严格按客户原件模板（邓老师 2026-06-08「完全按模板」）：
    周表只有 目的港 | 服务 | day1-7 | 备注，不含起运港列、不含币种列。
    起运港固定 PVG 不入表；币种不入表（回流由 AirAdapter 默认补 PVG/CNY）。"""
    rows = [{
        "origin": "PVG", "destination": "NRT", "service": "CA",
        "currency": "JPY", "day1": 10, "day2": 11,
        "effective_week_start": "2026-05-25",
    }]
    content, _ = fill_template("air", rows)  # 无 tier_prices → 走周表分支
    ws = load_workbook(BytesIO(content)).active
    header = [c.value for c in ws[1]]
    assert "Currency" not in header
    assert "Origin (POL)" not in header
    assert header[0] == "Destinations"  # A 列即目的港
    assert ws.cell(2, 1).value == "NRT"
