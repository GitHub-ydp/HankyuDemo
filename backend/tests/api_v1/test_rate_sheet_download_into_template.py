"""指定数据下载端点：multipart(模板 + rows JSON) → 回填后的 xlsx。"""
import json
from io import BytesIO
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.api.deps import get_db
from app.main import app

FIXTURE = (
    Path(__file__).resolve().parents[1] / "fixtures" / "sea_other_ports_template.xlsx"
)


@pytest.fixture
def client():
    def _override_get_db():
        yield None

    app.dependency_overrides[get_db] = _override_get_db
    yield TestClient(app)
    app.dependency_overrides.pop(get_db, None)


def _new_sea_session(client) -> str:
    r = client.post("/api/v1/rate-sheet/session", data={"template_type": "sea"})
    return r.json()["data"]["session_id"]


def test_download_into_template_fills_other_ports(client):
    sid = _new_sea_session(client)
    rows = [{"destination": "BUSAN", "carrier": "EAS",
             "container_20gp": 160, "container_40hq": 320}]
    r = client.post(
        f"/api/v1/rate-sheet/{sid}/download-into-template",
        data={"rows": json.dumps(rows)},
        files={"template": ("tpl.xlsx", FIXTURE.read_bytes(),
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert r.status_code == 200
    assert r.content[:2] == b"PK"
    ws = load_workbook(BytesIO(r.content))["FCL N RATE OF OTHER PORTS"]
    assert ws.cell(9, 1).value == "HONG KONG"   # 未匹配港留名
    # BUSAN 被填（在表里某行）
    a = [ws.cell(rr, 1).value for rr in range(9, ws.max_row + 1)]
    assert "BUSAN" in a


def test_large_rows_field_over_1mb_accepted(client):
    """整本合约的上万行运价作为单个 multipart 字段提交（>1MB）应被受理。

    Starlette 默认单字段上限 1MB（max_part_size），上万行 rows JSON 会超限被框架
    以 HTTP 400 "Part exceeded maximum size of 1024KB." 拒收。回归此场景。
    """
    sid = _new_sea_session(client)
    rows = [{"destination": "BUSAN", "carrier": "EAS",
             "container_20gp": 160, "container_40hq": 320,
             "remark": "padding " * 8} for _ in range(8000)]
    payload = json.dumps(rows)
    assert len(payload.encode("utf-8")) > 1024 * 1024  # 确认确实超 1MB
    r = client.post(
        f"/api/v1/rate-sheet/{sid}/download-into-template",
        data={"rows": payload},
        files={"template": ("tpl.xlsx", FIXTURE.read_bytes(),
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert r.status_code == 200
    assert r.content[:2] == b"PK"


def test_unknown_session_returns_404(client):
    r = client.post(
        "/api/v1/rate-sheet/nope/download-into-template",
        data={"rows": "[]"},
        files={"template": ("tpl.xlsx", FIXTURE.read_bytes(), "application/octet-stream")},
    )
    assert r.json()["code"] == 404


def test_bad_rows_json_returns_400(client):
    sid = _new_sea_session(client)
    r = client.post(
        f"/api/v1/rate-sheet/{sid}/download-into-template",
        data={"rows": "not-json"},
        files={"template": ("tpl.xlsx", FIXTURE.read_bytes(), "application/octet-stream")},
    )
    assert r.json()["code"] == 400


def test_non_xlsx_template_returns_400(client):
    sid = _new_sea_session(client)
    r = client.post(
        f"/api/v1/rate-sheet/{sid}/download-into-template",
        data={"rows": "[]"},
        files={"template": ("x.txt", b"not a zip", "text/plain")},
    )
    assert r.json()["code"] == 400


def test_air_session_returns_400(client):
    r = client.post("/api/v1/rate-sheet/session", data={"template_type": "air"})
    sid = r.json()["data"]["session_id"]
    r = client.post(
        f"/api/v1/rate-sheet/{sid}/download-into-template",
        data={"rows": "[]"},
        files={"template": ("tpl.xlsx", FIXTURE.read_bytes(), "application/octet-stream")},
    )
    assert r.json()["code"] == 400
