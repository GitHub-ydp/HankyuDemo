from __future__ import annotations
import re
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path
from openpyxl import load_workbook


# 整名别名：招标书写法 → 价源/字典规范名（词级别名走 rate_parser.PORT_ALIAS_MAP）
_POD_ALIASES = {
    "TANJUNG PRIOK": "JAKARTA",   # 雅加达港区名，价源按 JAKARTA 报价
}


def normalize_pod(raw: str) -> str:
    s = (raw or "").strip().upper()
    s = s.replace("（", "(").replace("）", ")")  # 全角括号 → 半角（招标书混用）
    s = s.replace("KELANG", "KLANG")          # 拼写统一 KELANG->KLANG
    s = s.split("(")[0].strip()               # 去括号注解，留主名
    return _POD_ALIASES.get(s, s)


@dataclass(slots=True)
class CostLane:
    pol: str
    pod_raw: str
    pod_norm: str
    carrier: str
    rate_20gp: Decimal | None
    rate_40hc: Decimal | None
    lss: str
    transit_time: str
    no_service: bool


def _dec(v):
    if v is None:
        return None
    try:
        return Decimal(str(v))
    except Exception:
        return None


def _first_int(v):
    m = re.search(r"\d+", str(v or ""))
    return int(m.group()) if m else None


class NitoriCostBook:
    def __init__(self, lanes: list[CostLane], origin_fees: dict, free_time: dict):
        self._lanes = lanes
        self.origin_fees = origin_fees
        self.free_time = free_time
        self._index = {(l.pol, l.pod_norm): l for l in lanes}

    @classmethod
    def from_xlsx(cls, path: Path) -> "NitoriCostBook":
        wb = load_workbook(path, data_only=True)
        ws = wb["FCL"]
        lanes: list[CostLane] = []
        for r in range(2, ws.max_row + 1):
            pol = ws.cell(r, 1).value
            dest = ws.cell(r, 2).value
            carrier = ws.cell(r, 3).value
            if not pol or not dest:
                continue
            pol_s = str(pol).strip().upper()
            if pol_s not in ("SHANGHAI", "TAICANG"):
                continue
            no_service = str(carrier).strip().upper() == "NO SERVICE"
            lanes.append(CostLane(
                pol=pol_s, pod_raw=str(dest), pod_norm=normalize_pod(str(dest)),
                carrier="" if no_service else str(carrier).strip(),
                rate_20gp=_dec(ws.cell(r, 4).value) if not no_service else None,
                rate_40hc=_dec(ws.cell(r, 5).value) if not no_service else None,
                lss=str(ws.cell(r, 8).value or ""),
                transit_time=str(ws.cell(r, 11).value or ""),
                no_service=no_service,
            ))
        free_time = cls._parse_free_time(wb)
        wb.close()
        return cls(lanes, origin_fees={}, free_time=free_time)

    def lookup(self, *, pol: str, pod: str, carrier: str | None = None) -> CostLane | None:
        return self._index.get((pol.strip().upper(), normalize_pod(pod)))

    def free_time_for(self, pod: str) -> dict | None:
        return self.free_time.get(normalize_pod(pod))

    @staticmethod
    def _parse_free_time(wb) -> dict:
        out = {}
        if "Sheet1" in wb.sheetnames:
            ws = wb["Sheet1"]
            for r in range(2, ws.max_row + 1):
                port = ws.cell(r, 1).value
                if not port:
                    continue
                out[normalize_pod(str(port))] = {
                    "dem": _first_int(ws.cell(r, 2).value),
                    "det": _first_int(ws.cell(r, 3).value),
                }
        return out
