from __future__ import annotations
from pathlib import Path
import extract_msg
from openpyxl import load_workbook


def _has_fcl_sheet(path: Path) -> bool:
    try:
        wb = load_workbook(path, read_only=True)
        try:
            return "FCL" in wb.sheetnames
        finally:
            wb.close()
    except Exception:
        return False


def resolve_bundle(folder: Path) -> tuple[Path, Path]:
    """从已解压的 Nitori 投标包目录定位 (TO GLOBAL 报价表, 成本表)。

    成本表识别依据：含名为 'FCL' 的 sheet（区别于 ① 邮件夹带的旧报价/合同 xlsx）。
    """
    folder = Path(folder)
    # 注意：不能用 next(...) 裸取——找不到时抛 StopIteration，在 async 端点里会被
    # Python 转成 RuntimeError 裸 500（无 CORS 头→前端误报 Network Error）。显式判空。
    quotes = [p for p in folder.glob("*GLOBAL*.xlsm") if not p.name.startswith("._")]
    if not quotes:
        raise FileNotFoundError(
            "TO GLOBAL 报价表(*GLOBAL*.xlsm) 未在投标包中找到"
        )
    quote = quotes[0]

    # 候选成本 xlsx：顶层散落 + 所有 .msg 抽出的 xlsx 附件
    candidates: list[Path] = list(folder.glob("*.xlsx"))
    for msg_path in folder.glob("*.msg"):
        m = extract_msg.Message(str(msg_path))
        try:
            for a in m.attachments:
                name = a.longFilename or a.shortFilename or ""
                if name.lower().endswith(".xlsx"):
                    out = folder / name
                    if not out.exists():
                        out.write_bytes(a.data)
                    candidates.append(out)
        finally:
            m.close()

    for c in candidates:
        if _has_fcl_sheet(c):
            return quote, c
    raise FileNotFoundError("成本表(含 FCL sheet 的 xlsx) 未在投标包中找到")
