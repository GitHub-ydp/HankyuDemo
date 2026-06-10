from __future__ import annotations
import shutil
from datetime import date as _date
from decimal import Decimal
from pathlib import Path
from openpyxl import load_workbook
from app.services.step1_rates.writers.base import is_formula_cell, safe_set, stamp_document_properties
from app.services.step2_bidding.entities import (
    ParsedPkg, PkgRow, PkgSection, PerRowReport, RowStatus, CostType,
)
from app.services.step2_bidding.nitori_cost_book import NitoriCostBook, normalize_pod

# ---- 常量（业务可调）----
_QUOTE_SHEET = "Quotation (Global) Jul-Sep"
_DATA_START_ROW = 9
_COL = dict(carrier=4, country_exp=5, pol=6, country_imp=7, pod=8, size=9,
            tt=17, of_cur=18, of_amt=19, lss_cur=20, lss_amt=21,
            thc_cur=26, thc_amt=27, doc_cur=28, doc_amt=29, dem_free=103, det_free=104)
_CHINA_POLS = {"SHANGHAI", "TAICANG"}
_SIZE_TO_COST = {"20F": "20gp", "40HC": "40hc", "40F": "40hc"}   # 40F←40HC（assumption）
_SIZE_TO_CONTAINER = {"20F": "container_20gp", "40HC": "container_40hq", "40F": "container_40hq"}  # 40HC=40HQ
_MARKUP = Decimal("1.15")


def _select_ocean_candidate(cands, col):
    """选价策略(邓老师口径 2026-06-10)：最新期间优先，同期内取最低价。

    返回 (chosen, pool)：chosen=选中行(无可用行时 None)；pool=同最新期间的全部
    可用候选(供标注复核提示)。期间以 valid_to 为准(缺失回退 valid_from，再缺
    视为最旧)，仅在该箱型有价的行中比较。
    """
    viable = [c for c in cands if getattr(c, col, None) is not None]
    if not viable:
        return None, []

    def _period(c):
        return c.valid_to or c.valid_from or _date.min

    latest = max(_period(c) for c in viable)
    pool = [c for c in viable if _period(c) == latest]
    chosen = min(pool, key=lambda c: getattr(c, col))
    return chosen, pool


class NitoriProfile:
    customer_code = "nitori"
    display_name = "ニトリ (Nitori) TO GLOBAL"
    priority = 20

    def __init__(self, cost_book: NitoriCostBook | None = None,
                 markup_ratio: Decimal = _MARKUP, repo=None):
        self._cost = cost_book
        self._markup = markup_ratio
        self._repo = repo

    def detect(self, path: Path, hint: str | None = None) -> bool:
        if hint == self.customer_code:
            return True
        try:
            wb = load_workbook(path, data_only=True, read_only=True)
        except Exception:
            return False
        try:
            return _QUOTE_SHEET in wb.sheetnames
        finally:
            wb.close()

    def parse(self, path: Path, bid_id: str, period: str) -> ParsedPkg:
        wb = load_workbook(path, data_only=True)
        try:
            ws = wb[_QUOTE_SHEET]
            rows: list[PkgRow] = []
            for r in range(_DATA_START_ROW, ws.max_row + 1):
                pol = ws.cell(r, _COL["pol"]).value
                pod = ws.cell(r, _COL["pod"]).value
                size = ws.cell(r, _COL["size"]).value
                if not pol or not pod:
                    continue
                pol_s = str(pol).strip().upper()
                rows.append(PkgRow(
                    row_idx=r, section_index=0, section_code="GLOBAL",
                    origin_code=pol_s, origin_text_raw=str(pol),
                    destination_text_raw=str(pod), destination_code=normalize_pod(str(pod)),
                    cost_type=CostType.UNKNOWN, currency="USD",
                    volume_desc=None, existing_price=None, existing_lead_time=None,
                    existing_carrier=str(ws.cell(r, _COL["carrier"]).value or "") or None,
                    existing_remark=None, is_example=False, client_constraint_text=None,
                    extras={"size": str(size).strip() if size else "",
                            "pod_raw": str(pod), "is_china": pol_s in _CHINA_POLS},
                ))
            section = PkgSection(0, "GLOBAL", 6, "CHINA", "CN", "USD", "", False, [])
            return ParsedPkg(bid_id=bid_id, customer_code=self.customer_code, period=period,
                             sheet_name=_QUOTE_SHEET, source_file=path.name,
                             sections=[section], rows=rows, warnings=[])
        finally:
            wb.close()

    def match(self, parsed: ParsedPkg) -> list[PerRowReport]:
        reports: list[PerRowReport] = []
        for row in parsed.rows:
            if not row.extras.get("is_china"):
                continue
            size = row.extras.get("size", "")
            hit = self._match_from_db(row, size)
            if hit is None and self._cost is not None:
                hit = self._match_from_cost_book(row, size)
            if hit is None:
                reports.append(PerRowReport(
                    row_idx=row.row_idx, section_code="GLOBAL",
                    destination_code=row.destination_code, status=RowStatus.NO_RATE,
                    cost_price=None, sell_price=None, markup_ratio=None,
                    lead_time_text=None, carrier_text=None,
                    remark_text="库内无该航线运价，待询价",
                    selected_candidate=None))
                continue
            sell = (hit["price"] * self._markup).quantize(Decimal("1"))
            reports.append(PerRowReport(
                row_idx=row.row_idx, section_code="GLOBAL",
                destination_code=row.destination_code, status=RowStatus.FILLED,
                cost_price=hit["price"], sell_price=sell, markup_ratio=self._markup,
                lead_time_text=hit["lead"], carrier_text=hit["carrier"],
                remark_text=hit["note"], selected_candidate=None,
                thc_amount=hit["thc"], doc_amount=hit["doc"], lss_amount=hit["lss"]))
        return reports

    def _match_from_db(self, row, size):
        """优先：查 DB 运价(福山确认口径)。命中返回 dict，否则 None。

        选价规则(邓老师口径 2026-06-10)：先取最新期间(valid_to 最大)，
        同一期间内多个报价取最低价，并在 note 标注候选明细提醒复核。
        """
        if self._repo is None:
            return None
        col = _SIZE_TO_CONTAINER.get(size)
        if not col:
            return None
        cands = self._repo.query_ocean_fcl(origin=row.origin_code, destination=row.destination_code)
        chosen, pool = _select_ocean_candidate(cands, col)
        if chosen is None:
            return None
        price = getattr(chosen, col)
        lead = chosen.transit_time_text or (str(chosen.transit_days) if chosen.transit_days is not None else None)
        note = None
        if len(pool) > 1:
            others = "、".join(
                f"{c.carrier_name or '?'} {getattr(c, col)}" for c in pool if c is not chosen
            )
            note = (
                f"同航线同期 {len(pool)} 个报价，已取最低价 {chosen.carrier_name or '?'} {price}"
                f"（其余：{others}），请复核"
            )
        is_20 = size == "20F"
        lss = (chosen.lss_20 if is_20 else chosen.lss_40) or chosen.lss_cic
        return {
            "price": price, "carrier": chosen.carrier_name, "lead": lead,
            "thc": None if is_20 else chosen.thc,   # 库中 thc 经 20/40 行合并≈40 箱档值，20F 宁缺勿错
            "doc": chosen.doc, "lss": lss, "note": note,
        }

    def _match_from_cost_book(self, row, size):
        """回退：zip 内成本文件(MVP 兜底；福山验收后或移除)。命中返回 dict，否则 None。"""
        lane = self._cost.lookup(pol=row.origin_code, pod=row.destination_code)
        cost_field = _SIZE_TO_COST.get(size)
        if not lane or lane.no_service or not cost_field:
            return None
        price = getattr(lane, f"rate_{cost_field}")
        if price is None:
            return None
        return {"price": price, "carrier": lane.carrier, "lead": lane.transit_time,
                "thc": None, "doc": None, "lss": None, "note": None}

    def fill(self, source_path: Path, parsed: ParsedPkg,
             row_reports: list[PerRowReport], variant: str, output_path: Path):
        if variant not in ("cost", "sr"):
            raise ValueError(f"variant 必须 cost/sr，实际 {variant!r}")
        shutil.copy2(source_path, output_path)
        wb = load_workbook(output_path, data_only=False, keep_vba=True)   # 保宏
        try:
            ws = wb[_QUOTE_SHEET]
            for rep in row_reports:
                if rep.status != RowStatus.FILLED:
                    continue
                price = rep.cost_price if variant == "cost" else rep.sell_price
                self._set(ws, rep.row_idx, "carrier", rep.carrier_text)
                self._set(ws, rep.row_idx, "of_cur", "USD")
                self._set(ws, rep.row_idx, "of_amt", float(price))
                if rep.lss_amount is not None:                    # 未知不写，不再硬编 0
                    self._set(ws, rep.row_idx, "lss_cur", "USD")
                    self._set(ws, rep.row_idx, "lss_amt", float(rep.lss_amount))
                if rep.thc_amount is not None:
                    self._set(ws, rep.row_idx, "thc_cur", "CNY")
                    self._set(ws, rep.row_idx, "thc_amt", float(rep.thc_amount))
                if rep.doc_amount is not None:
                    self._set(ws, rep.row_idx, "doc_cur", "CNY")
                    self._set(ws, rep.row_idx, "doc_amt", float(rep.doc_amount))
                if rep.lead_time_text:
                    self._set(ws, rep.row_idx, "tt", rep.lead_time_text)
                ft = self._cost.free_time_for(rep.destination_code) if self._cost else None
                if ft:
                    self._set(ws, rep.row_idx, "dem_free", ft.get("dem"))
                    self._set(ws, rep.row_idx, "det_free", ft.get("det"))
            stamp_document_properties(wb, batch_id=f"{parsed.bid_id}:nitori:{variant}")
            wb.save(output_path)
        finally:
            wb.close()

    @staticmethod
    def _set(ws, row_idx: int, col_key: str, value):
        if value is None:
            return
        cell = ws.cell(row_idx, _COL[col_key])
        if is_formula_cell(cell):
            return
        safe_set(cell, value)
