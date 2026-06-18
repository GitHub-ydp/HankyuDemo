"""配置驱动的模板填充器。

输入统一的 normalized rate dict 列表，按 template_registry 的列映射，从数据起始行
逐行写进空白模板（保留表头/格式/公式）。复用 writers/base.safe_set 的写入守卫
（None 不写、公式格不覆盖）。不入库——产物就是填好的 xlsx。

空白模板的数据区预设了合并单元格（如目的港/船司列每 2 行合并），其从属格 value 只读。
填充前先解除「数据起始行及以后」的合并（表头区合并保留），再逐格写。

normalized rate dict 字段约定：
  通用:  destination, carrier, remark
  sea :  container_20gp, container_40gp, container_40hq, lss_cic, baf, ebs, yas_caf, sailing, via, transit, booking
  air(周表源 Market Price):  service, day1..day7
  air(档位源 EES/唯凯):       service, tier_prices(稀疏 KG→价)

档位源(EES/唯凯)有多个重量档。默认「严格按模板」：只取 +100KG 一档(模板 C 列就是
Service/+100KG)，套进 air_blank.xlsx 周表(价铺满 day1-7)。关掉 strict 才走旧的「程序从零
生成动态档位表(起运港|目的港|服务|动态 KG 列|备注)」——见 _STRICT_AIR_TEMPLATE_DEFAULT。
"""
from __future__ import annotations

from copy import copy
from datetime import date, timedelta
from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook

from app.services.step1_rates.sheet_builder.entities import SheetFillConfig
from app.services.step1_rates.sheet_builder.template_registry import get_template_config
from app.services.step1_rates.writers.base import safe_set

# Sea FCL 一条运价展开为三行：(箱型标签, 取运费用的字段名)
_SEA_CONTAINER_ROWS = (
    ("20FT", "container_20gp"),
    ("40GP", "container_40gp"),
    ("40HQ", "container_40hq"),
)

# (表头标签, 列键, 行字段) —— 新增海运元数据列(模板无表头,填充时一并写第8行表头)
_SEA_META_COLS = (
    ("Currency", "currency", "currency"),
    ("Valid From", "valid_from", "valid_from"),
    ("Valid To", "valid_to", "valid_to"),
    ("Rate Level", "rate_level", "rate_level"),
    ("Service Code", "service_code", "service_code"),
)

# 档位表元数据列：(表头标签, 取值的字段名)。顺序即列序，必须与 AirTierAdapter 解析契约一致。
_TIER_META_COLS = (
    ("Currency", "currency"),
    ("Effective From", "effective_week_start"),
    ("Effective To", "effective_to"),
    ("Carrier", "carrier"),
    ("Cargo Class", "cargo_class"),
    ("Packing", "packing"),
    ("Density", "density"),
)

# air 档位源默认出表方式：True=严格套 air_blank.xlsx，只取 +100KG 档(邓老师 2026-06-08
# 要求「完全按模板」)；False=动态多档表(_build_tier_sheet，旧行为，保留备查/可切回)。
# 客户确认后若要回到多档，改这一处常量(或调用 fill_template 时传 strict_air_template=False)即可。
_STRICT_AIR_TEMPLATE_DEFAULT = True


def fill_template(
    template_type: str,
    rows: list[dict[str, Any]],
    *,
    strict_air_template: bool | None = None,
) -> tuple[bytes, str]:
    """把 rows 填进对应空白模板，返回 (xlsx_bytes, 建议文件名)。

    air 档位源(行带 tier_prices)有两种出表方式，由 strict_air_template 决定
    (None → 取模块默认 _STRICT_AIR_TEMPLATE_DEFAULT)：
      - True(默认)：严格套 air_blank.xlsx 周表模板——每条航线只取 +100KG 一档
        (模板 C 列 = Service/+100KG)，价铺满 day1-7；其余档位(45/300/500/1000…)按模板舍弃。
      - False：走 _build_tier_sheet 动态多档表(列=全表档位并集)，旧行为保留备查。
    """
    cfg = get_template_config(template_type)  # 校验类型；未知抛 ValueError

    # air 档位源(行带 tier_prices)：默认严格按模板(降档为 +100KG 周表行，下落 _fill_air)；
    # 关掉 strict 则回到动态档位表(_build_tier_sheet，旧码保留)。
    if template_type == "air" and any(r.get("tier_prices") for r in rows):
        strict = (
            _STRICT_AIR_TEMPLATE_DEFAULT
            if strict_air_template is None
            else strict_air_template
        )
        if not strict:
            return _build_tier_sheet(rows)
        rows = _tier_rows_as_weekly(rows)

    workbook = load_workbook(cfg.template_path, data_only=False)
    if template_type == "air":
        _fill_air(workbook, cfg.sheets[0], rows)
    elif template_type == "sea":
        _fill_sea(workbook, cfg.sheets[0], rows)
    else:  # pragma: no cover — registry 已保证只有 air/sea
        raise ValueError(f"unsupported template_type {template_type!r}")

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue(), f"{template_type}_rate_sheet_filled.xlsx"


def _row_tiers(row: dict[str, Any]) -> dict[int, Any]:
    """取行的档位 dict，键归一为 int(KG)——JSON 往返后键可能是字符串('45')。"""
    raw = row.get("tier_prices") or {}
    return {int(kg): price for kg, price in raw.items()}


def _tier_sheet_name(rows: list[dict[str, Any]]) -> str:
    """报价日(effective_week_start)入 sheet 名；无则用通用名。"""
    for row in rows:
        value = row.get("effective_week_start")
        if value:
            return f"Air Rates {str(value)[:10]}"
    return "Air Tier Rates"


def _build_tier_sheet(rows: list[dict[str, Any]]) -> tuple[bytes, str]:
    """程序生成档位表：起运港|目的港|服务|动态 KG 列|元数据列|备注。"""
    tiers = sorted({kg for row in rows for kg in _row_tiers(row)})
    wb = Workbook()
    ws = wb.active
    ws.title = _tier_sheet_name(rows)

    header = (
        ["Origin (POL)", "Destination", "Service"]
        + [f"{kg}KG" for kg in tiers]
        + [label for label, _ in _TIER_META_COLS]
        + ["Remark"]
    )
    for c, label in enumerate(header, start=1):
        ws.cell(1, c).value = label

    meta_start = 4 + len(tiers)  # KG 列之后第一列
    r = 2
    for row in rows:
        ws.cell(r, 1).value = row.get("origin")
        ws.cell(r, 2).value = row.get("destination")
        ws.cell(r, 3).value = row.get("service")
        row_tiers = _row_tiers(row)
        for i, kg in enumerate(tiers):
            price = row_tiers.get(kg)
            if price is not None:
                ws.cell(r, 4 + i).value = price
        for j, (_, field_name) in enumerate(_TIER_META_COLS):
            ws.cell(r, meta_start + j).value = row.get(field_name)
        ws.cell(r, meta_start + len(_TIER_META_COLS)).value = row.get("remark")
        r += 1

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue(), "air_tier_rate_sheet_filled.xlsx"


def _pick_p100(tiers: dict[int, float]) -> tuple[float | None, int | None]:
    """从稀疏档位取 +100KG 代表价。精确 100 优先；无则取最接近 100 的档
    (距离相等取较小档)。返回 (价, 实际取的档位 KG)；空档返回 (None, None)。"""
    if not tiers:
        return None, None
    if 100 in tiers:
        return tiers[100], 100
    kg = min(tiers, key=lambda k: (abs(k - 100), k))
    return tiers[kg], kg


def _tier_rows_as_weekly(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """严格按模板：把档位行降为 air_blank.xlsx 周表行——只留 +100KG 一档，
    铺满 day1-7(档位价无每日维度，整周同价)。已是周表行(无 tier_prices)原样透传，
    兼容同批混排。非 100 档代表时在备注标注实际档位，便于审核追溯。"""
    out: list[dict[str, Any]] = []
    for row in rows:
        tiers = _row_tiers(row)
        if not tiers:
            out.append(row)
            continue
        price, used_kg = _pick_p100(tiers)
        new = {k: v for k, v in row.items() if k != "tier_prices"}
        for day in range(1, 8):
            new[f"day{day}"] = price
        if used_kg is not None and used_kg != 100:
            note = f"(按 {used_kg}KG 档)"
            new["remark"] = f"{(row.get('remark') or '').strip()} {note}".strip()
        out.append(new)
    return out


def _unmerge_data_area(ws, data_start_row: int) -> None:
    """解除「数据起始行及以后」的合并单元格，让数据区每格可写。表头合并保留。"""
    targets = [
        str(rng) for rng in ws.merged_cells.ranges if rng.min_row >= data_start_row
    ]
    for rng in targets:
        ws.unmerge_cells(rng)


def _first_week_start(rows: list[dict[str, Any]]) -> date | None:
    """取第一条带 effective_week_start 的行的周起始日（ISO 字符串 / date 均可）；无则 None。"""
    for row in rows:
        value = row.get("effective_week_start")
        if not value:
            continue
        try:
            return date.fromisoformat(str(value)[:10])
        except ValueError:
            return None
    return None


def _apply_week_headers(ws, sheet_cfg: SheetFillConfig, rows: list[dict[str, Any]]) -> None:
    """有周信息时，按该周改写 day1-7 日期表头 + sheet 名（一份模板通吃任意周）；无则不动。"""
    week_start = _first_week_start(rows)
    if week_start is None:
        return
    col = sheet_cfg.columns
    for i in range(7):
        d = week_start + timedelta(days=i)
        ws.cell(sheet_cfg.header_row, col[f"day{i + 1}"]).value = (
            f"{d.year}/{d.month}/{d.day} ({d:%a})"
        )
    week_end = week_start + timedelta(days=6)
    ws.title = f"{week_start:%b} {week_start.day} to {week_end:%b} {week_end.day}"


def _fill_air(workbook, sheet_cfg: SheetFillConfig, rows: list[dict[str, Any]]) -> None:
    # 严格按客户原件布局：目的港 | 服务 | day1-7 | 备注。无起运港/币种列——
    # 起运港固定 PVG 不入表(回流由 AirAdapter 默认补 PVG/CNY)；row 里的 origin/currency 不写。
    ws = workbook[sheet_cfg.sheet_name]
    _unmerge_data_area(ws, sheet_cfg.data_start_row)
    _apply_week_headers(ws, sheet_cfg, rows)
    col = sheet_cfg.columns
    r = sheet_cfg.data_start_row
    for row in rows:
        safe_set(ws.cell(r, col["destination"]), row.get("destination"))
        safe_set(ws.cell(r, col["service"]), row.get("service"))
        for day in range(1, 8):
            safe_set(ws.cell(r, col[f"day{day}"]), row.get(f"day{day}"))
        safe_set(ws.cell(r, col["remark"]), row.get("remark"))
        r += 1


def _fill_sea(workbook, sheet_cfg: SheetFillConfig, rows: list[dict[str, Any]]) -> None:
    ws = workbook[sheet_cfg.sheet_name]
    _unmerge_data_area(ws, sheet_cfg.data_start_row)
    col = sheet_cfg.columns
    # 模板无这些新列表头 → 填充时在表头行写英文标签，供重新导入时 OceanAdapter 按表头识别
    for label, col_key, _ in _SEA_META_COLS:
        ws.cell(sheet_cfg.header_row, col[col_key]).value = label
    r = sheet_cfg.data_start_row
    for row in rows:
        for container_label, freight_key in _SEA_CONTAINER_ROWS:
            safe_set(ws.cell(r, col["destination"]), row.get("destination"))
            safe_set(ws.cell(r, col["carrier"]), row.get("carrier"))
            safe_set(ws.cell(r, col["container"]), container_label)
            safe_set(ws.cell(r, col["freight"]), row.get(freight_key))
            safe_set(ws.cell(r, col["lss_cic"]), row.get("lss_cic"))
            safe_set(ws.cell(r, col["baf"]), row.get("baf"))
            safe_set(ws.cell(r, col["ebs"]), row.get("ebs"))
            safe_set(ws.cell(r, col["yas_caf"]), row.get("yas_caf"))
            safe_set(ws.cell(r, col["sailing"]), row.get("sailing"))
            safe_set(ws.cell(r, col["via"]), row.get("via"))
            transit_cell = ws.cell(r, col["transit"])
            safe_set(transit_cell, row.get("transit"))
            # 模板 Transit Time 列按行做了纵向合并、且偶数行藏着美元/时间数字格式。解除合并后
            # 仅设 number_format='General' 在存盘时会被合并样式覆盖(openpyxl 已知坑)，导致天数
            # 渲染成「$6」。用 style='Normal' 彻底重置该格样式(=General)，再从同行运费列(始终
            # 带表格边框)补回边框，确保 transit 显示纯数字「6」且不丢表格线。
            if transit_cell.value is not None:
                kept_border = copy(ws.cell(r, col["freight"]).border)
                transit_cell.style = "Normal"
                transit_cell.border = kept_border
            safe_set(ws.cell(r, col["booking"]), row.get("booking"))
            safe_set(ws.cell(r, col["rmks"]), row.get("remark"))
            for _, col_key, field_name in _SEA_META_COLS:
                safe_set(ws.cell(r, col[col_key]), row.get(field_name))
            r += 1
