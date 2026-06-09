"""按客户上传的模板回填海运运价（仅 FCL N RATE OF OTHER PORTS 页）。

与 template_filler.fill_template（顺序填内置空白模板）并存、互不影响。
本模块：读上传 workbook → 扫 A 列保留的目的港 → 按 canonicalize 匹配当前会话 rows →
以数据为准重排「船司 × {20FT, 40FT/40HQ}」行（抓原模板样式重铺、保留版式）→ 返回 bytes。
对不上/无数据的港留空（不报告）。其余 sheet（JP/LCL）原样保留。
"""
from __future__ import annotations

import re
from copy import copy
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from app.services.step1_rates.port_normalizer import canonicalize
from app.services.step1_rates.writers.base import safe_set, save_workbook_to_bytes


class RefillError(Exception):
    """模板不满足回填前提（如缺目标工作表）。端点据此返回 400。"""


# 本次只配 OTHER PORTS 一份；以后扩 JP/LCL 只加 profile，不改算法。
OTHER_PORTS_PROFILE: dict[str, Any] = {
    "sheet_name": "FCL N RATE OF OTHER PORTS",
    "header_row": 8,
    "data_start_row": 9,
    "cols": {
        "destination": 1, "carrier": 2, "container": 3, "freight": 4,
        "lss": 5, "baf": 6, "cic": 7, "caf": 8,
        "sailing": 9, "via": 10, "transit": 11,
        "booking": 12, "thc": 13, "doc": 14, "isps": 15, "equipment": 16,
        "rmks": 17,
    },
    "container_rows": [
        {"label": "20FT", "freight": ["container_20gp"], "container": 20},
        {"label": "40FT/40HQ", "freight": ["container_40hq", "container_40gp"], "container": 40},
    ],
    "surcharge_cols": {"lss": "LSS", "baf": "BAF", "cic": "CIC", "caf": "CAF"},
    # 每船司块内跨 2 行合并的列（复刻原模板：船司/船期/中转/航程/备注竖向合并）
    "merge_cols": ("carrier", "sailing", "via", "transit", "rmks"),
}

_SPLIT_RE = re.compile(r"[/\n]")


def refill_into_template(
    template_bytes: bytes,
    rows: list[dict[str, Any]],
    *,
    profile: dict[str, Any] = OTHER_PORTS_PROFILE,
) -> bytes:
    """把 rows 回填进上传模板的目标页，返回填好的 xlsx bytes。

    非 .xlsx/损坏 → openpyxl 抛异常（端点转 400）；缺目标页 → RefillError。
    """
    wb = load_workbook(BytesIO(template_bytes), data_only=False)
    name = profile["sheet_name"]
    if name not in wb.sheetnames:
        raise RefillError(f"模板缺少 '{name}' 工作表")
    ws = wb[name]

    ports = _scan_ports(ws, profile)
    styles = _capture_row_styles(ws, profile)
    _clear_data_area(ws, profile)
    by_port = _group_rows_by_port(rows)
    _write_ports(ws, profile, ports, by_port, styles)
    return save_workbook_to_bytes(wb)


def _scan_ports(ws, profile: dict[str, Any]) -> list[str]:
    """从数据起始行往下扫 A 列，非空即一个目的港，按出现顺序返回。"""
    col = profile["cols"]["destination"]
    ports: list[str] = []
    for r in range(profile["data_start_row"], ws.max_row + 1):
        v = ws.cell(r, col).value
        if v is not None and str(v).strip() != "":
            ports.append(str(v).strip())
    return ports


def _capture_row_styles(ws, profile: dict[str, Any]) -> dict[str, dict[int, Any]]:
    """抓数据区前两行（20FT 行 / 40 行）每列样式，作为重铺新行的样式模板。"""
    start = profile["data_start_row"]
    maxc = ws.max_column
    return {
        "c20": {c: copy(ws.cell(start, c)._style) for c in range(1, maxc + 1)},
        "c40": {c: copy(ws.cell(start + 1, c)._style) for c in range(1, maxc + 1)},
    }


def _clear_data_area(ws, profile: dict[str, Any]) -> None:
    """unmerge 数据区所有合并格并清值（样式已被 _capture_row_styles 抓走）。

    该模板目的港后无页脚（已确认），故清到 max_row。
    """
    start = profile["data_start_row"]
    for rng in [str(m) for m in ws.merged_cells.ranges if m.min_row >= start]:
        ws.unmerge_cells(rng)
    for r in range(start, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            ws.cell(r, c).value = None


def _group_rows_by_port(rows: list[dict[str, Any]]) -> dict[str, list[dict[str, Any]]]:
    out: dict[str, list[dict[str, Any]]] = {}
    for row in rows or []:
        key = canonicalize(row.get("destination"))
        if not key:
            continue
        out.setdefault(key, []).append(row)
    return out


def _port_candidates(name: str) -> set[str]:
    """模板港名 → 候选 canonical 集合：去括号、按 / 与换行拆名，提升命中率。

    例：'MADRAS / CHENNAI'→{MADRAS,CHENNAI,MADRASCHENNAI}；
        'CHICAGO (via LAX)'→{CHICAGO}；'LONG BEACH\\nLOS ANGELES'→{LONGBEACH,LOSANGELES,...}
    """
    base = re.sub(r"\([^)]*\)", " ", str(name or ""))
    cands: set[str] = set()
    whole = canonicalize(base)
    if whole:
        cands.add(whole)
    for part in _SPLIT_RE.split(base):
        cp = canonicalize(part)
        if cp:
            cands.add(cp)
    return cands


def _match_rows(port_name: str, by_port: dict[str, list[dict[str, Any]]]) -> list[dict[str, Any]]:
    matched: list[dict[str, Any]] = []
    for cand in _port_candidates(port_name):
        matched.extend(by_port.get(cand, []))
    return matched


def _freight(row: dict[str, Any], crow: dict[str, Any]) -> Any:
    for key in crow["freight"]:
        v = row.get(key)
        if v is not None:
            return v
    return None


def _surcharge_cell(row: dict[str, Any], code: str, container: int) -> Any:
    """结构化附加费 → 单元格值，还原原始样本写法（Incl./Collect/数值/备注）。"""
    amt_key = "amount_20" if container == 20 else "amount_40"
    for item in (row.get("surcharges") or []):
        if str(item.get("code") or "").strip().upper() != code:
            continue
        if item.get("included"):
            return "Incl."
        if str(item.get("payment") or "").strip().lower() == "collect":
            return "Collect"
        amt = item.get(amt_key)
        if amt is None:
            amt = item.get("amount")
        if amt is not None:
            return amt
        if item.get("note"):
            return item.get("note")
        return None
    # 兜底：surcharges 无此项 → 用扁平字段（老 Excel 行）
    if code == "LSS":
        return row.get("lss_cic")
    if code == "BAF":
        return row.get("baf")
    return None


def _apply_style(ws, r: int, style_map: dict[int, Any]) -> None:
    for c, st in style_map.items():
        ws.cell(r, c)._style = copy(st)


def _write_ports(ws, profile, ports, by_port, styles) -> None:
    col = profile["cols"]
    crows = profile["container_rows"]
    scol = profile["surcharge_cols"]
    # 有数据的港按模板序紧凑排在前（动态行），无数据的港留名追加在后。
    matched_ports = [(p, m) for p in ports if (m := _match_rows(p, by_port))]
    empty_ports = [p for p in ports if not _match_rows(p, by_port)]
    r = profile["data_start_row"]
    for port, matched in matched_ports:
        block_start = r
        for row in matched:
            top, bot = crows[0], crows[1]
            # 20FT 行
            _apply_style(ws, r, styles["c20"])
            safe_set(ws.cell(r, col["destination"]), port)
            safe_set(ws.cell(r, col["carrier"]), row.get("carrier"))
            ws.cell(r, col["container"]).value = top["label"]
            safe_set(ws.cell(r, col["freight"]), _freight(row, top))
            safe_set(ws.cell(r, col["lss"]), _surcharge_cell(row, scol["lss"], 20))
            safe_set(ws.cell(r, col["baf"]), _surcharge_cell(row, scol["baf"], 20))
            safe_set(ws.cell(r, col["cic"]), _surcharge_cell(row, scol["cic"], 20))
            safe_set(ws.cell(r, col["caf"]), _surcharge_cell(row, scol["caf"], 20))
            safe_set(ws.cell(r, col["via"]), row.get("via"))
            safe_set(ws.cell(r, col["transit"]), row.get("transit") or row.get("transit_days"))
            safe_set(ws.cell(r, col["rmks"]), row.get("remark"))
            # 40FT/40HQ 行
            _apply_style(ws, r + 1, styles["c40"])
            ws.cell(r + 1, col["container"]).value = bot["label"]
            safe_set(ws.cell(r + 1, col["freight"]), _freight(row, bot))
            safe_set(ws.cell(r + 1, col["lss"]), _surcharge_cell(row, scol["lss"], 40))
            safe_set(ws.cell(r + 1, col["baf"]), _surcharge_cell(row, scol["baf"], 40))
            safe_set(ws.cell(r + 1, col["cic"]), _surcharge_cell(row, scol["cic"], 40))
            safe_set(ws.cell(r + 1, col["caf"]), _surcharge_cell(row, scol["caf"], 40))
            # 每船司块跨 2 行合并（复刻原模板竖向合并）
            for key in profile["merge_cols"]:
                ws.merge_cells(start_row=r, end_row=r + 1,
                               start_column=col[key], end_column=col[key])
            r += 2
        # 港名（A 列）竖向合并整块
        if r - 1 > block_start:
            ws.merge_cells(start_row=block_start, end_row=r - 1,
                           start_column=col["destination"], end_column=col["destination"])
    # 无数据的港：保留港名 + 1 空行（让她看到这个港没报到价）
    for port in empty_ports:
        _apply_style(ws, r, styles["c20"])
        safe_set(ws.cell(r, col["destination"]), port)
        r += 1
