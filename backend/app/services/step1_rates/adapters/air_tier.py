"""Step1 空运重量档(air_tier)导入适配器。

两种来料都入 AirTierRate：
1. 「做表」生成的档位表（英文契约表头 Origin/Destination/Service/{kg}KG）；
2. 客户原始 EES 风格中文档位表（目的港/航班/比重/≧45/100/500/1000KG，
   复用做表侧 air_ees.parse_ees——合并单元格前向填充/区域多港/议价排除等规则一致）。
按 sheet 表头内容识别（不靠文件名——档位文件名常含 'air' 会被 AirAdapter 抢）。
"""
from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.services.step1_rates.entities import (
    ParsedRateBatch,
    ParsedRateRecord,
    Step1FileType,
)
from app.services.step1_rates.sheet_builder.air_ees import (
    is_ees_header_row,
    parse_ees,
)

_TIER_HEADER_RE = re.compile(r"^(\d+)\s*KG$", re.IGNORECASE)
_EXCEL_EXTS = {".xlsx", ".xlsm", ".xls"}
# EES 表头(目的港/≧100KG)通常在 sheet 前 10 行内(上方是公告/生效说明)，扫 40 行留足余量
_EES_HEADER_SCAN_ROWS = 40


class AirTierAdapter:
    """识别并解析做表生成的空运重量档表。"""

    key = "air_tier"
    file_type = Step1FileType.air_tier
    priority = 3  # 独占低值，先于 nvo_fak(5)/air_weekly(6)/air(10)

    def detect(self, path: Path, *, file_type_hint: Step1FileType | None = None) -> bool:
        if file_type_hint is not None:
            return file_type_hint == self.file_type
        if path.suffix.lower() not in _EXCEL_EXTS:
            return False
        try:
            wb = load_workbook(path, data_only=True, read_only=True)
        except Exception:
            return False
        try:
            for ws in wb.worksheets:
                if self._tier_sheet_headers(ws) is not None:
                    return True
                if self._has_ees_header(ws):
                    return True
        finally:
            wb.close()
        return False

    def parse(self, path: Path, db: Session | None = None) -> ParsedRateBatch:
        wb = load_workbook(path, data_only=True)
        records: list[ParsedRateRecord] = []
        warnings: list[str] = []
        try:
            matched = False
            for ws in wb.worksheets:
                headers = self._tier_sheet_headers(ws)
                if headers is None:
                    continue
                matched = True
                records.extend(self._parse_sheet(ws, headers))
        finally:
            wb.close()
        if not matched:
            # 英文契约未命中 → 试 EES 中文档位表(目的港/航班/比重/≧45…KG)
            records, ees_warnings, matched = self._parse_ees(path)
            warnings.extend(ees_warnings)
        if not matched:
            warnings.append(
                "未找到符合档位表表头契约的 sheet"
                "(Origin/Destination/Service/{kg}KG 或 目的港/≧100KG…)"
            )
        return ParsedRateBatch(
            file_type=Step1FileType.air_tier,
            source_file=path.name,
            records=records,
            warnings=warnings,
            adapter_key=self.key,
        )

    @staticmethod
    def _has_ees_header(ws) -> bool:
        for row in ws.iter_rows(
            min_row=1, max_row=_EES_HEADER_SCAN_ROWS, values_only=True
        ):
            if is_ees_header_row(list(row)):
                return True
        return False

    def _parse_ees(
        self, path: Path
    ) -> tuple[list[ParsedRateRecord], list[str], bool]:
        """EES 中文档位表 → ParsedRateRecord 列表。返回 (records, warnings, matched)。

        matched 以「解析出行」为准：EES 表头判定在 parse_ees 内部完成，
        没有任何航线行说明该文件并非 EES 档位表。
        """
        result = parse_ees(str(path))
        rows: list[dict[str, Any]] = result.get("parsed_rows") or []
        if not rows:
            return [], [], False
        records = [
            ParsedRateRecord(
                record_kind="air_tier",
                # EES 为上海起运，起运港留空由 mapper 默认 PVG（与做表口径一致）
                origin_port_name=None,
                destination_port_name=row.get("destination_port_name"),
                service_desc=row.get("service_desc"),
                currency="CNY",  # EES 全表人民币含油价
                valid_from=row.get("effective_week_start"),
                valid_to=None,
                remarks=row.get("remarks"),
                source_type="excel",
                extras={
                    "tier_prices": row.get("tier_prices") or {},
                    "carrier": row.get("carrier"),
                    "row_index": idx,
                },
            )
            for idx, row in enumerate(rows, start=1)
        ]
        return records, list(result.get("warnings") or []), True

    def _tier_sheet_headers(self, ws) -> dict[str, Any] | None:
        """读第 1 行表头；命中返回 {"named": {表头小写→列下标}, "tiers": [(kg, 列下标)]}，否则 None。"""
        first_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if not first_row:
            return None
        named: dict[str, int] = {}
        tiers: list[tuple[int, int]] = []
        for idx, cell in enumerate(first_row):
            text = str(cell).strip() if cell is not None else ""
            if not text:
                continue
            m = _TIER_HEADER_RE.match(text)
            if m:
                tiers.append((int(m.group(1)), idx))
            else:
                named[text.lower()] = idx
        has_origin = any(k.startswith("origin") for k in named)
        if has_origin and "destination" in named and tiers:
            return {"named": named, "tiers": tiers}
        return None

    def _parse_sheet(self, ws, headers: dict[str, Any]) -> list[ParsedRateRecord]:
        named: dict[str, int] = headers["named"]
        tiers: list[tuple[int, int]] = headers["tiers"]

        def col(*names: str) -> int | None:
            for n in names:
                if n in named:
                    return named[n]
            return None

        c_origin = col("origin (pol)", "origin")
        c_dest = col("destination")
        c_service = col("service")
        c_currency = col("currency")
        c_from = col("effective from")
        c_to = col("effective to")
        c_carrier = col("carrier")
        c_cargo = col("cargo class")
        c_pack = col("packing")
        c_density = col("density")
        c_remark = col("remark")

        out: list[ParsedRateRecord] = []
        for row_index, row in enumerate(
            ws.iter_rows(min_row=2, values_only=True), start=2
        ):
            def v(i: int | None) -> Any:
                return row[i] if i is not None and i < len(row) else None

            origin = v(c_origin)
            dest = v(c_dest)
            tier_prices: dict[int, float] = {}
            for kg, ci in tiers:
                num = self._safe_float(row[ci] if ci < len(row) else None)
                if num is not None:
                    tier_prices[kg] = num
            if not origin and not dest and not tier_prices:
                continue  # 跳过空行
            out.append(
                ParsedRateRecord(
                    record_kind="air_tier",
                    origin_port_name=str(origin).strip() if origin else None,
                    destination_port_name=str(dest).strip() if dest else None,
                    service_desc=(str(v(c_service)).strip() if v(c_service) else None),
                    currency=(str(v(c_currency)).strip() if v(c_currency) else "CNY"),
                    valid_from=self._to_date(v(c_from)),
                    valid_to=self._to_date(v(c_to)),
                    remarks=(str(v(c_remark)).strip() if v(c_remark) else None),
                    source_type="excel",
                    extras={
                        "tier_prices": tier_prices,
                        "cargo_class": (str(v(c_cargo)).strip() if v(c_cargo) else None),
                        "packing": (str(v(c_pack)).strip() if v(c_pack) else None),
                        "density": (str(v(c_density)).strip() if v(c_density) else None),
                        "carrier": (str(v(c_carrier)).strip() if v(c_carrier) else None),
                        "row_index": row_index,
                    },
                )
            )
        return out

    @staticmethod
    def _to_date(value: Any) -> date | None:
        if not value:
            return None
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        text = str(value).strip()[:10]
        for fmt in ("%Y-%m-%d", "%Y/%m/%d"):
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                continue
        return None

    @staticmethod
    def _safe_float(value: Any) -> float | None:
        """数字格→float；空/非数字(如 'ASK'/'-'/'询价')返回 None，单格脏数据不炸整批。"""
        if value is None or str(value).strip() == "":
            return None
        try:
            return float(value)
        except (TypeError, ValueError):
            return None
