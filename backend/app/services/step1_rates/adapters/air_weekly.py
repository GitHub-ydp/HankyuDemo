"""Step1 空运周报(air_weekly)回流适配器。

读「做表」生成的周报表布局(air_blank)：
  Origin (POL) | Destinations | Service/+100KG | <7 个日期列> | Remark (Selling) | Currency
产 record_kind="air_weekly" → activator 已接线 air→AirFreightRate。
按表头内容识别(不靠文件名)；周起始日从首个日期列表头(形如 '2026/5/25 (Mon)')解析。
注意与既有 AirAdapter(读真实承运商布局 A1='Destinations'、无 origin 列)区分：本适配器要求有 'origin' 表头。
"""
from __future__ import annotations

import re
from datetime import date, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.services.step1_rates.entities import (
    ParsedRateBatch,
    ParsedRateRecord,
    Step1FileType,
)

_DATE_RE = re.compile(r"(\d{4})/(\d{1,2})/(\d{1,2})")
_EXCEL_EXTS = {".xlsx", ".xlsm", ".xls"}


class AirWeeklyAdapter:
    """识别并解析做表生成的空运周报表(air_blank 布局)。"""

    key = "air_weekly"
    file_type = Step1FileType.air
    priority = 6  # air_tier(5) 之后、air(10) 之前

    def detect(self, path: Path, *, file_type_hint: Step1FileType | None = None) -> bool:
        if file_type_hint is not None:
            return False  # air hint 仍交给真实承运商 AirAdapter
        if path.suffix.lower() not in _EXCEL_EXTS:
            return False
        try:
            wb = load_workbook(path, data_only=True, read_only=True)
        except Exception:
            return False
        try:
            for ws in wb.worksheets:
                if self._weekly_headers(ws) is not None:
                    return True
        finally:
            wb.close()
        return False

    def parse(self, path: Path, db: Session | None = None) -> ParsedRateBatch:
        wb = load_workbook(path, data_only=True)
        records: list[ParsedRateRecord] = []
        eff_from: date | None = None
        eff_to: date | None = None
        for ws in wb.worksheets:
            headers = self._weekly_headers(ws)
            if headers is None:
                continue
            if eff_from is None and headers.get("week_start"):
                eff_from = headers["week_start"]
                eff_to = headers["week_start"] + timedelta(days=6)
            records.extend(self._parse_sheet(ws, headers))
        return ParsedRateBatch(
            file_type=Step1FileType.air,
            source_file=path.name,
            effective_from=eff_from,
            effective_to=eff_to,
            records=records,
            adapter_key=self.key,
        )

    def _weekly_headers(self, ws) -> dict[str, Any] | None:
        """命中周报布局返回 {origin,destination,service,remark,currency,date_cols:[idx],week_start}，否则 None。"""
        first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if not first:
            return None
        named: dict[str, int] = {}
        date_cols: list[int] = []
        week_start: date | None = None
        for idx, cell in enumerate(first):
            text = str(cell).strip() if cell is not None else ""
            if not text:
                continue
            m = _DATE_RE.search(text)
            if m:
                date_cols.append(idx)
                if week_start is None:
                    week_start = date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
            else:
                named[text.lower()] = idx
        has_origin = any(k.startswith("origin") for k in named)
        has_dest = any(k.startswith("destination") for k in named)
        has_service = any(k.startswith("service") for k in named)
        if has_origin and has_dest and has_service and len(date_cols) >= 7:
            return {
                "origin": next(named[k] for k in named if k.startswith("origin")),
                "destination": next(named[k] for k in named if k.startswith("destination")),
                "service": next(named[k] for k in named if k.startswith("service")),
                "remark": next((named[k] for k in named if k.startswith("remark")), None),
                "currency": named.get("currency"),
                "date_cols": date_cols[:7],
                "week_start": week_start,
            }
        return None

    def _parse_sheet(self, ws, h: dict[str, Any]) -> list[ParsedRateRecord]:
        week_start: date | None = h["week_start"]
        week_end = week_start + timedelta(days=6) if week_start else None
        out: list[ParsedRateRecord] = []
        for row_index, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            def v(i: int | None) -> Any:
                return row[i] if i is not None and i < len(row) else None

            origin = v(h["origin"])
            dest = v(h["destination"])
            prices = [self._dec(row[ci] if ci < len(row) else None) for ci in h["date_cols"]]
            if not origin and not dest and not any(p is not None for p in prices):
                continue
            kw: dict[str, Any] = {f"price_day{i + 1}": prices[i] for i in range(7)}
            out.append(
                ParsedRateRecord(
                    record_kind="air_weekly",
                    origin_port_name=str(origin).strip() if origin else None,
                    destination_port_name=str(dest).strip() if dest else None,
                    service_desc=(str(v(h["service"])).strip() if v(h["service"]) else None),
                    currency=(str(v(h["currency"])).strip() if v(h["currency"]) else "CNY"),
                    effective_week_start=week_start,
                    effective_week_end=week_end,
                    remarks=(str(v(h["remark"])).strip() if v(h["remark"]) else None),
                    source_type="excel",
                    extras={"row_index": row_index},
                    **kw,
                )
            )
        return out

    @staticmethod
    def _dec(value: Any) -> Decimal | None:
        if value is None or (isinstance(value, str) and value.strip() == ""):
            return None
        try:
            return value if isinstance(value, Decimal) else Decimal(str(value))
        except (InvalidOperation, ValueError):
            return None
